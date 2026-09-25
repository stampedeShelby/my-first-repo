"""Small python-docx helpers used by build_report.py (keeps the template's look:
Times New Roman, justified body, bordered tables, numbered figures/tables)."""

from __future__ import annotations

import copy
import re

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

FONT = "Times New Roman"
MONO = "Courier New"
BULLET_NUM_ID = 90
TEXT_WIDTH_DXA = 9026  # A4 minus 1" margins


def _set_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


class Report:
    def __init__(self, doc):
        self.doc = doc
        self.fig_no = {}
        self.tab_no = {}
        self.chapter = 0
        self._ensure_styles()
        self._ensure_bullets()

    # ------------------------------------------------------------------ setup
    def _ensure_styles(self):
        styles = self.doc.styles
        spec = {"Heading 1": (16, 0, 18, 12), "Heading 2": (13, 1, 12, 6), "Heading 3": (12, 2, 8, 4)}
        for name, (size, lvl, before, after) in spec.items():
            try:
                st = styles[name]
            except KeyError:
                st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = styles["Normal"]
            st.font.name = FONT
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = RGBColor(0, 0, 0)
            pf = st.paragraph_format
            pf.space_before, pf.space_after = Pt(before), Pt(after)
            pf.keep_with_next = True
            ppr = st.element.get_or_add_pPr()
            for old in ppr.findall(qn("w:outlineLvl")):
                ppr.remove(old)
            ol = OxmlElement("w:outlineLvl")
            ol.set(qn("w:val"), str(lvl))
            ppr.append(ol)
            rpr = st.element.get_or_add_rPr()
            rf = OxmlElement("w:rFonts")
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(a), FONT)
            rpr.insert(0, rf)

    def _ensure_bullets(self):
        numbering = self.doc.part.numbering_part.element
        abs_id = "90"
        absn = OxmlElement("w:abstractNum")
        absn.set(qn("w:abstractNumId"), abs_id)
        mlt = OxmlElement("w:multiLevelType"); mlt.set(qn("w:val"), "hybridMultilevel"); absn.append(mlt)
        for lvl, (char, font, ind) in enumerate([("•", "Symbol", 720), ("o", "Courier New", 1440)]):
            l = OxmlElement("w:lvl"); l.set(qn("w:ilvl"), str(lvl))
            for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", "" if lvl == 0 else "o"),
                             ("w:lvlJc", "left")):
                e = OxmlElement(tag); e.set(qn("w:val"), val); l.append(e)
            ppr = OxmlElement("w:pPr"); indent = OxmlElement("w:ind")
            indent.set(qn("w:left"), str(ind)); indent.set(qn("w:hanging"), "360"); ppr.append(indent); l.append(ppr)
            rpr = OxmlElement("w:rPr"); rf = OxmlElement("w:rFonts")
            for a in ("w:ascii", "w:hAnsi", "w:hint"):
                rf.set(qn(a), font if a != "w:hint" else "default")
            rpr.append(rf); l.append(rpr)
            absn.append(l)
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(absn)
        else:
            numbering.append(absn)
        num = OxmlElement("w:num"); num.set(qn("w:numId"), str(BULLET_NUM_ID))
        a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), abs_id); num.append(a)
        numbering.append(num)

    # ------------------------------------------------------------------ text
    def _runs(self, p, text, size=12, base_bold=False, color=None):
        for part in re.split(r"(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)", text):
            if not part:
                continue
            if part.startswith("**"):
                _set_font(p.add_run(part[2:-2]), size=size, bold=True, color=color)
            elif part.startswith("`"):
                _set_font(p.add_run(part[1:-1]), name=MONO, size=size - 1.5, bold=base_bold, color=color)
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                _set_font(p.add_run(part[1:-1]), size=size, italic=True, bold=base_bold, color=color)
            else:
                _set_font(p.add_run(part), size=size, bold=base_bold, color=color)

    def _fmt(self, p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, line=1.5, before=0):
        pf = p.paragraph_format
        pf.alignment = align
        pf.space_after = Pt(after)
        pf.space_before = Pt(before)
        pf.line_spacing = line
        return p

    def para(self, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, after=6, bold=False):
        p = self._fmt(self.doc.add_paragraph(), align=align, after=after)
        self._runs(p, text, size=size, base_bold=bold)
        return p

    def bullets(self, items, level=0, size=12):
        for it in items:
            p = self._fmt(self.doc.add_paragraph(), after=2, line=1.3)
            ppr = p._p.get_or_add_pPr()
            numpr = OxmlElement("w:numPr")
            il = OxmlElement("w:ilvl"); il.set(qn("w:val"), str(level)); numpr.append(il)
            ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(BULLET_NUM_ID)); numpr.append(ni)
            ppr.append(numpr)
            self._runs(p, it, size=size)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(8)

    def h1(self, title, first=False):
        self.chapter += 1
        self.fig_no[self.chapter] = 0
        self.tab_no[self.chapter] = 0
        p = self.doc.add_paragraph(style="Heading 1")
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.page_break_before = not first
        p.paragraph_format.space_after = Pt(18)
        _set_font(p.add_run(f"CHAPTER {self.chapter}"), size=14, bold=True)
        p.add_run().add_break(WD_BREAK.LINE)
        _set_font(p.add_run(title.upper()), size=16, bold=True)
        return p

    def h2(self, title):
        p = self.doc.add_paragraph(style="Heading 2")
        _set_font(p.add_run(title), size=13, bold=True)
        return p

    def h3(self, title):
        p = self.doc.add_paragraph(style="Heading 3")
        _set_font(p.add_run(title), size=12, bold=True)
        return p

    def code(self, text, size=8.5):
        p = self._fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.LEFT, after=8, line=1.0)
        p.paragraph_format.left_indent = Cm(0.3)
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F4F7")
        ppr.append(shd)
        bdr = OxmlElement("w:pBdr")
        for side in ("top", "left", "bottom", "right"):
            b = OxmlElement(f"w:{side}"); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "4"); b.set(qn("w:color"), "C9CED6"); bdr.append(b)
        ppr.insert(0, bdr)
        lines = text.rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            r = p.add_run(line)
            _set_font(r, name=MONO, size=size)
            if i < len(lines) - 1:
                r.add_break(WD_BREAK.LINE)
        return p

    # ------------------------------------------------------------------ figures / tables
    def figure(self, path, caption, width_cm=15.5):
        p = self._fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, after=2, line=1.0, before=6)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        self.fig_no[self.chapter] += 1
        c = self._fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, after=10, line=1.0)
        _set_font(c.add_run(f"Figure {self.chapter}.{self.fig_no[self.chapter]}: "), size=10.5, bold=True)
        _set_font(c.add_run(caption), size=10.5, italic=True)
        return f"Figure {self.chapter}.{self.fig_no[self.chapter]}"

    def table(self, header, rows, widths_cm, caption=None, size=10, header_fill="D9E2F3", bold_first_col=False):
        if caption:
            self.tab_no[self.chapter] += 1
            c = self._fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, after=4, line=1.0, before=6)
            c.paragraph_format.keep_with_next = True
            _set_font(c.add_run(f"Table {self.chapter}.{self.tab_no[self.chapter]}: "), size=10.5, bold=True)
            _set_font(c.add_run(caption), size=10.5, italic=True)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblpr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblpr.append(lay)
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths_cm[i] * 567)))
        for ri, row in enumerate([header] + rows):
            tr = t.rows[ri]
            if ri == 0:
                trpr = tr._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
            cant = OxmlElement("w:cantSplit"); tr._tr.get_or_add_trPr().append(cant)
            for ci, val in enumerate(row):
                cell = tr.cells[ci]
                cell.width = Cm(widths_cm[ci])
                p = cell.paragraphs[0]
                self._fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, after=1, line=1.05)
                self._runs(p, str(val), size=size, base_bold=(ri == 0 or (bold_first_col and ci == 0)))
                if ri == 0:
                    tcpr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), header_fill); tcpr.append(shd)
        spacer = self._fmt(self.doc.add_paragraph(), after=4, line=1.0)
        return t
