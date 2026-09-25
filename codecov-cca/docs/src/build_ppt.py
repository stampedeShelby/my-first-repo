"""Fills the BCS701 PPT template with the CCA presentation.

python docs/src/build_ppt.py <template.pptx> <out.pptx>
"""
import copy
import json
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "..", "figures")
RESULTS = json.load(open(os.path.join(HERE, "..", "..", "demo", "results.json")))

NAVY = RGBColor(0x1B, 0x2F, 0x5B)
RED = RGBColor(0xB3, 0x26, 0x1E)
GREEN = RGBColor(0x1E, 0x7D, 0x46)
INK = RGBColor(0x1D, 0x24, 0x33)
MUTED = RGBColor(0x4A, 0x52, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xC9, 0xD1, 0xE0)
TINT = {"navy": RGBColor(0xEE, 0xF1, 0xFB), "red": RGBColor(0xFD, 0xEC, 0xEA),
        "green": RGBColor(0xE6, 0xF4, 0xEC), "white": WHITE}
SERIF, SANS = "Times New Roman", "Calibri"

prs = Presentation(sys.argv[1])
TEMPLATE_BLANK = prs.slides[2]
footer_parts = [copy.deepcopy(sh._element) for sh in TEMPLATE_BLANK.shapes if sh.is_placeholder]
layout = TEMPLATE_BLANK.slide_layout

# drop the two empty template slides; we add our own after the Content slide
sldIdLst = prs.slides._sldIdLst
for sid in list(sldIdLst)[2:]:
    prs.part.drop_rel(sid.rId)
    sldIdLst.remove(sid)


# ------------------------------------------------------------------ helpers
def new_slide(title, notes):
    s = prs.slides.add_slide(layout)
    for el in list(s.shapes._spTree):
        if el.tag.endswith("}sp") and el.find(".//{*}ph") is not None:
            s.shapes._spTree.remove(el)
    for el in footer_parts:
        s.shapes._spTree.append(copy.deepcopy(el))
    tb = s.shapes.add_textbox(Inches(0.45), Inches(0.78), Inches(9.1), Inches(0.62))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    run = tf.paragraphs[0].add_run()
    run.text = title
    style(run, 28, True, NAVY, SERIF)
    s.notes_slide.notes_text_frame.text = notes
    return s


def style(run, size, bold=False, color=INK, font=SANS, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = color


def text(s, x, y, w, h, parts, size=16, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         font=SANS, bullets=False, spacing=4):
    """parts: list of paragraphs; each paragraph a str or list of (text, bold) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    for i, para in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.25)))
            pPr.set("indent", str(-Inches(0.22)))
            bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
            pPr.append(bu)
        runs = [(para, False)] if isinstance(para, str) else para
        for t, b in runs:
            r = p.add_run()
            r.text = t
            style(r, size, b, color, font)
    return tb


def card(s, x, y, w, h, tint="white", border=LINE, radius=True):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = TINT[tint]
    shp.line.color.rgb = border
    shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def circle(s, x, y, d, label, color=NAVY, size=14):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = color
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    style(r, size, True, WHITE)


def picture(s, name, x, y, w=None, h=None, border=True):
    path = os.path.join(FIG, name)
    iw, ih = Image.open(path).size
    if w and h:  # fit inside box, centred
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        x, y = x + (w - pw) / 2, y + (h - ph) / 2
    elif w:
        pw, ph = w, w * ih / iw
    else:
        pw, ph = h * iw / ih, h
    if border:
        card(s, x - 0.06, y - 0.06, pw + 0.12, ph + 0.12, radius=False)
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(pw), Inches(ph))
    return y + ph


def table(s, x, y, w, col_w, rows, size=12, header_fill=NAVY, row_h=0.36, highlight=None):
    shape = s.shapes.add_table(len(rows), len(col_w), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = shape.table
    tblPr = t._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblPr.set(attr, "0")
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for r_i, row in enumerate(rows):
        t.rows[r_i].height = Inches(row_h)
        for c_i, val in enumerate(row):
            cell = t.cell(r_i, c_i)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            color = INK
            if r_i == 0:
                cell.fill.fore_color.rgb = header_fill
                color = WHITE
            else:
                cell.fill.fore_color.rgb = WHITE if r_i % 2 else TINT["navy"]
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run(); r.text = val
            bold = r_i == 0
            if highlight and r_i > 0:
                hc = highlight(r_i, c_i, val)
                if hc is not None:
                    color, bold = hc, True
            style(r, size, bold, color)
    return shape


# ------------------------------------------------------------------ slide 1 (title)
s1 = prs.slides[0]
for sh in s1.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text
    if t.startswith("Application Development"):
        runs = sh.text_frame.paragraphs[0].runs
        runs[0].text = "CCA: Security Breach Analysis"
        runs[1].text = ""
        runs[2].text = ""
        for r in runs:
            r.font.size = Pt(28)
    elif t.startswith("Application Title"):
        sh.left, sh.width = Inches(0.4), Inches(9.2)
        sh.top = Inches(2.35)
        sh.text_frame.word_wrap = True
        p = sh.text_frame.paragraphs[0]
        p.runs[0].text = "Application Title: "
        p.runs[1].text = "SignGate"
        p.runs[1].font.bold = True
        p.runs[1].font.color.rgb = NAVY
        p2 = sh.text_frame.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r = p2.add_run()
        r.text = "Stopping the Codecov (2021) Supply-Chain Attack with a Cryptographic Verify-Before-Execute Gate"
        style(r, 18, False, INK, SERIF, italic=True)
    elif t.startswith("Presented By"):
        sh.top = Inches(3.55)
    elif t.startswith("Under the guidance"):
        sh.top = Inches(5.35)
text(s1, 2.0, 3.95, 6.0, 1.2,
     ["<Student Name 1>  –  <USN>", "<Student Name 2>  –  <USN>", "<Student Name 3>  –  <USN>"],
     size=15, align=PP_ALIGN.CENTER, font=SERIF, spacing=2)
text(s1, 2.0, 5.72, 6.0, 0.75, [[("Prof. Sonnegowda K", True)], "Assistant Professor, Dept. of ISE, BMSIT&M"],
     size=15, align=PP_ALIGN.CENTER, font=SERIF, spacing=0)
s1.notes_slide.notes_text_frame.text = (
    "Good morning. Our CCA case study is the 2021 Codecov supply-chain attack. We will show how one "
    "leaked key let an attacker steal secrets from thousands of CI pipelines, and present SignGate, a "
    "small cryptographic gate we built that would have stopped it. (~20 s)")

# ------------------------------------------------------------------ slide 2 (contents)
s2 = prs.slides[1]
box = [sh for sh in s2.shapes if sh.name == "TextBox 6"][0]
items = ["Title of the application", "Abstract", "Introduction: The Codecov Incident", "Problem Statement",
         "Objectives", "Requirements", "Methodology", "Design and Implementation",
         "Security Analysis", "Results (System Testing)", "Innovation and Relevance",
         "Conclusion", "References : (IEEE format)"]
paras = box.text_frame.paragraphs
template_p = paras[1]._p
for p in list(paras)[1:]:
    p._p.getparent().remove(p._p)
paras[0].runs[0].text = items[0]
for item in items[1:]:
    new_p = copy.deepcopy(template_p)
    box.text_frame._txBody.append(new_p)
    runs = new_p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}r")
    runs[0].find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = item
for p in box.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(20)
    epr = p._p.find("{http://schemas.openxmlformats.org/drawingml/2006/main}endParaRPr")
    if epr is not None:
        epr.set("sz", "2000")
box.top = Inches(1.55)
s2.notes_slide.notes_text_frame.text = "This is the flow of our presentation. (~5 s)"

# ------------------------------------------------------------------ 3 abstract
s = new_slide("Abstract",
    "In one line: an attacker modified a trusted tool that thousands of companies ran automatically, "
    "and nobody noticed for two months. We analysed why this worked and built SignGate, which signs "
    "every release and verifies it inside CI before running anything. We tested it against seven "
    "attack scenarios and all malicious ones were blocked. (~30 s)")
text(s, 0.5, 1.55, 9.0, 3.2, [
    [("Incident: ", True), ("In 2021 an HMAC key leaked from a Docker image let an attacker modify the "
      "Codecov Bash Uploader. Thousands of CI/CD pipelines ran it and leaked their secrets for about 2 months.", False)],
    [("Solution: ", True), ("SignGate, a verify-before-execute gate. SHA-256 + Ed25519 signed releases, "
      "a root-certified key hierarchy with revocation, a hash-chained transparency log, anti-rollback, "
      "and least-privilege execution.", False)],
    [("Result: ", True), ("Python prototype blocked the real attack and 5 stronger variants, "
      "cut exposed secrets from 8 to 1, and adds only ~0.6 ms per build.", False)]],
    size=17, bullets=True, spacing=10)
stats = [("~2 months", "attack went undetected", RED, "red"), ("6 checks", "before any code runs", NAVY, "navy"),
         ("8 → 1", "CI secrets exposed", GREEN, "green"), (f"{RESULTS['verify_ms']} ms", "verification cost", NAVY, "navy")]
for i, (big, small, col, tint) in enumerate(stats):
    x = 0.5 + i * 2.28
    card(s, x, 4.85, 2.1, 1.85, tint)
    text(s, x, 5.05, 2.1, 0.8, [[(big, True)]], size=28, color=col, align=PP_ALIGN.CENTER)
    text(s, x, 5.85, 2.1, 0.6, [small], size=13, color=MUTED, align=PP_ALIGN.CENTER)

# ------------------------------------------------------------------ 4 introduction
s = new_slide("Introduction: The Codecov Incident",
    "Codecov is a code-coverage SaaS. Customers ran its Bash Uploader inside CI with curl piped to "
    "bash, so it ran next to all their secrets. From 31 January 2021 the attacker kept modifying it. "
    "On 1 April a customer noticed the SHA-256 did not match. The main violation is integrity and "
    "authenticity, and confidentiality was lost as a result. (~40 s)")
facts = [("Organisation", "Codecov: SaaS code-coverage tool, 29,000+ customers"),
         ("Sector", "DevOps / CI-CD, software supply chain"),
         ("Date", "31 Jan to 1 Apr 2021 (disclosed 15 Apr)")]
for i, (k, v) in enumerate(facts):
    x = 0.5 + i * 3.05
    card(s, x, 1.55, 2.9, 0.95, "navy")
    text(s, x + 0.1, 1.6, 2.7, 0.9, [[(k, True)], v], size=12.5, spacing=1)
picture(s, "timeline.png", 0.55, 2.75, w=8.9)
text(s, 0.5, 5.75, 3.0, 0.4, [[("Security requirements violated:", True)]], size=14)
pills = [("Integrity (primary)", RED, "red"), ("Authentication", RED, "red"), ("Confidentiality", RED, "red"),
         ("Key management", NAVY, "navy")]
for i, (label, col, tint) in enumerate(pills):
    x = 0.5 + i * 2.28
    card(s, x, 6.2, 2.12, 0.5, tint, border=col)
    text(s, x, 6.2, 2.12, 0.5, [[(label, True)]], size=13, color=col, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------------------ 5 problem statement
s = new_slide("Problem Statement",
    "Here is the attack chain. A GCS HMAC key was left inside a Docker image layer. The attacker used it "
    "to overwrite the uploader, and every pipeline that ran it sent env and git remote to the attacker. "
    "Six weaknesses made this possible: V1 to V6. Our problem: CI runs third-party code with full "
    "secrets but has no enforced way to verify it first. (~45 s)")
card(s, 0.5, 1.5, 9.0, 0.78, "red", border=RED)
text(s, 0.65, 1.5, 8.7, 0.78, [[("CI/CD pipelines download and execute third-party tools with full access to secrets, "
     "but have no enforced way to verify the tool is authentic and unmodified before running it.", True)]],
     size=14, color=RED, anchor=MSO_ANCHOR.MIDDLE)
picture(s, "attack.png", 0.55, 2.5, w=8.9)
vulns = ["V1  Secret left in public Docker image layer", "V2  Static, over-privileged storage HMAC key",
         "V3  Uploader not signed; SHA-256 not enforced", "V4  'curl | bash' runs whatever is served",
         "V5  Every CI step can read every secret", "V6  No monitoring → ~2 months dwell time"]
for i, v in enumerate(vulns):
    x = 0.5 + (i % 3) * 3.05
    y = 5.35 + (i // 3) * 0.72
    card(s, x, y, 2.9, 0.6, "white")
    text(s, x + 0.08, y, 2.78, 0.6, [[(v[:2], True), (v[2:], False)]], size=12, anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------------------ 6 objectives
s = new_slide("Objectives",
    "Six objectives, each mapped to a weakness. Integrity and authenticity are the core. Key management "
    "and misuse detection handle a stronger attacker who steals a key. Least privilege limits damage. "
    "And it must be simple and cheap enough to be adopted. (~25 s)")
objs = [("Integrity", "Detect any modification of the uploader before it executes"),
        ("Authenticity", "Accept only software signed by the genuine publisher key, even if the bucket and hash are attacker-controlled"),
        ("Key management", "Protected keys, 90-day rotation, instant revocation without changing customer pipelines"),
        ("Misuse & rollback", "Expose silent use of a stolen key; block downgrade to old vulnerable versions"),
        ("Least privilege", "If malicious code ever runs, it sees only the one token it needs"),
        ("Simple & cheap", "Standard algorithms, ~1 ms overhead, one-line CI change")]
for i, (h, d) in enumerate(objs):
    x = 0.5 + (i % 3) * 3.05
    y = 1.65 + (i // 3) * 2.6
    card(s, x, y, 2.9, 2.35, "white")
    circle(s, x + 0.2, y + 0.22, 0.55, f"O{i+1}", NAVY, 14)
    text(s, x + 0.2, y + 0.9, 2.55, 0.45, [[(h, True)]], size=17, color=NAVY)
    text(s, x + 0.2, y + 1.35, 2.55, 1.0, [d], size=13, color=INK)

# ------------------------------------------------------------------ 7 requirements
s = new_slide("Requirements",
    "Functional requirements describe what the system does: sign, certify, log, verify, restrict and "
    "monitor. Security requirements are measurable: any bit flip detected, unforgeable signatures, "
    "pinned trust anchor, protected keys, fail-closed, under 50 ms. The prototype needs only Python "
    "and one crypto library. (~30 s)")
cols = [("Functional", "navy", ["FR1 Sign release manifest", "FR2 Root certifies keys + revocation list",
                                 "FR3 Append every release to a transparency log", "FR4 Six-check verification in CI",
                                 "FR5 Execute with env allowlist", "FR6 Monitor served artifact"]),
        ("Security", "red", ["SR1 Detect any 1-bit change (SHA-256)", "SR2 Unforgeable signature (Ed25519)",
                             "SR3 Pinned trust anchor", "SR4 Keys in HSM/KMS, AES-256 at rest",
                             "SR5 Fail-closed", "SR6 < 50 ms overhead"]),
        ("Hardware / Software", "green", ["Any x86/ARM machine, 2 GB RAM", "Python 3.9+, pyca/cryptography",
                                          "Bash, Git, GitHub Actions", "Production: Cloud KMS / HSM",
                                          "Object storage / CDN over TLS 1.3"])]
for i, (h, tint, lines) in enumerate(cols):
    x = 0.5 + i * 3.05
    card(s, x, 1.6, 2.9, 4.6, tint)
    text(s, x + 0.15, 1.72, 2.6, 0.5, [[(h, True)]], size=19, color=NAVY)
    text(s, x + 0.1, 2.3, 2.7, 3.8, lines, size=15, bullets=True, spacing=10)

# ------------------------------------------------------------------ 8 methodology
s = new_slide("Methodology",
    "We followed six steps: study the incident from primary sources, find vulnerabilities and model "
    "threats with STRIDE, design cryptographic controls for each one, implement them in Python, "
    "attack our own implementation, and compare it with the standard fix. (~25 s)")
steps = [("Study incident", "Codecov security update, post-mortem, Reuters, HashiCorp bulletin"),
         ("Analyse", "Attack chain, vulnerabilities V1–V6, STRIDE threat model"),
         ("Design", "Map each weakness to a crypto / security control"),
         ("Implement", "Python + pyca/cryptography prototype (~500 lines)"),
         ("Attack & test", "7 attack scenarios, assume-breach test, 13 unit tests, timing"),
         ("Evaluate", "Compare with standard fix; cost and feasibility")]
for i, (h, d) in enumerate(steps):
    x = 0.5 + (i % 3) * 3.05
    y = 1.7 + (i // 3) * 2.5
    shp = s.shapes.add_shape(MSO_SHAPE.CHEVRON if i % 3 else MSO_SHAPE.PENTAGON, Inches(x), Inches(y),
                             Inches(2.9), Inches(0.7))
    shp.fill.solid(); shp.fill.fore_color.rgb = NAVY if i < 3 else GREEN
    shp.line.fill.background(); shp.shadow.inherit = False
    tf = shp.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run(); r.text = f"{i+1}. {h}"
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    style(r, 15, True, WHITE)
    card(s, x, y + 0.85, 2.9, 1.35, "white")
    text(s, x + 0.12, y + 0.92, 2.66, 1.25, [d], size=13.5)
text(s, 0.5, 6.55, 9.0, 0.4, [[("Tools: ", True), ("Python 3, pyca/cryptography (Ed25519, AES), Bash, unittest, GitHub Actions", False)]],
     size=13, color=MUTED)

# ------------------------------------------------------------------ 9 architecture
s = new_slide("Design: SignGate Architecture",
    "The key idea is to move trust from where the file is stored to who signed it. The publisher "
    "hashes and signs each release with a KMS-held key certified by an offline root and logs it. The "
    "bucket may be attacker-controlled; that no longer matters. In the customer's CI, SignGate runs "
    "six checks against keys pinned in the customer's own repo, then runs the tool with only the "
    "variables it needs. (~50 s)")
card(s, 0.5, 1.5, 9.0, 0.55, "navy", border=NAVY)
text(s, 0.5, 1.5, 9.0, 0.55, [[("Idea: move trust from the storage bucket to cryptographic keys", True)]],
     size=16, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
picture(s, "arch.png", 0.55, 2.25, w=8.9)

# ------------------------------------------------------------------ 10 crypto building blocks
s = new_slide("Design: Cryptographic Building Blocks",
    "These are the primitives and why each fits. SHA-256 detects change. Ed25519 proves who signed it, "
    "and it is fast and deterministic. The mini-PKI and revocation list handle key theft. The "
    "transparency log makes misuse visible. The important point: a hash alone is not enough, because "
    "whoever can change the file can change the published hash. A signature binds the hash to a key "
    "the attacker does not have. (~45 s)")
rows = [["Mechanism", "Standard", "Role in SignGate"],
        ["SHA-256", "FIPS 180-4", "Fingerprint of uploader; log hash chain"],
        ["Ed25519 signature", "RFC 8032 / FIPS 186-5", "Signs manifest, key certs, CRL, log head"],
        ["Mini-PKI (root → release key)", "NIST SP 800-57", "90-day certified keys; rotation"],
        ["Signed revocation list", "X.509-CRL style", "Blocks stolen keys / bad artifacts"],
        ["Hash-chained transparency log", "RFC 6962 (CT) idea", "Makes silent key misuse visible"],
        ["AES-256 PKCS#8, HSM/KMS", "FIPS 197", "Private keys never in plain text"],
        ["TLS 1.3 + workload identity", "RFC 8446, OIDC", "Transport; replaces static HMAC key"]]
table(s, 0.5, 1.55, 9.0, [2.9, 2.3, 3.8], rows, size=13, row_h=0.47)
card(s, 0.5, 5.5, 9.0, 1.2, "green", border=GREEN)
text(s, 0.7, 5.55, 8.6, 1.1, [[("Why a signature, not just a hash?  ", True),
     ("An attacker who can edit the file can also edit the published hash. "
      "σ = Sign(SK_release, SHA-256 manifest) can only be produced with the private key, and it is verified "
      "with a public key pinned in the customer's repo.", False)]], size=14, anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------------------ 11 key management
s = new_slide("Design: Key Management",
    "Key management is exactly where Codecov failed. Our root key is offline and only certifies "
    "release keys and signs the revocation list. Its public key is pinned by customers. Release keys "
    "live in KMS for 90 days and rotate automatically. If one leaks, the root signs a new CRL and "
    "every pipeline rejects it on the next run, with no customer changes. (~40 s)")
picture(s, "keys.png", 0.55, 1.6, w=8.9)
kb = [[("Separation of duties: ", True), ("root (offline HSM), release (KMS), log (separate operator)", False)],
      [("Rotation: ", True), ("new release key every 90 days; customers change nothing", False)],
      [("Revocation: ", True), ("root-signed CRL blocks a leaked key on every CI run immediately", False)],
      [("Root cause fixed: ", True), ("no long-lived secret inside build images; secret-scan every layer", False)]]
text(s, 0.5, 5.1, 9.0, 1.8, kb, size=15, bullets=True, spacing=5)

# ------------------------------------------------------------------ 12 implementation
s = new_slide("Implementation: Verify-Before-Execute Gate",
    "This is the gate's decision flow. The checks are ordered and fail-closed, so the first failure "
    "blocks execution. The code is split into small modules, as listed on the right. For a customer, "
    "adopting it is a one-line change: replace curl pipe bash with signgate run. (~45 s)")
picture(s, "flow.png", 0.55, 1.6, w=5.8, h=5.15)
mods = [["Module", "Does"], ["crypto_utils", "SHA-256, Ed25519, AES keys"], ["pki", "Key certs, revocation"],
        ["tlog", "Hash-chained log"], ["publisher", "Hash → sign → log → publish"],
        ["gate", "6 checks + env allowlist"], ["monitor", "Tamper alerts"]]
table(s, 6.6, 1.6, 2.9, [1.15, 1.75], mods, size=11, row_h=0.36)
card(s, 6.6, 4.35, 2.9, 0.95, "red", border=RED)
text(s, 6.68, 4.38, 2.75, 0.9, [[("Before", True)], "bash <(curl -s codecov.io/bash)"], size=11.5, color=RED, spacing=1)
card(s, 6.6, 5.45, 2.9, 1.3, "green", border=GREEN)
text(s, 6.68, 5.48, 2.75, 1.25, [[("After", True)], "python -m signgate.cli run --bucket b/ --policy policy.json"],
     size=11.5, color=GREEN, spacing=1)

# ------------------------------------------------------------------ 13 security analysis
s = new_slide("Security Analysis: STRIDE Threat Model",
    "In 2021 the attacker had bucket write access but no signing key. With SignGate, the edited file "
    "fails the hash check. Rewriting the hash breaks the signature. Using their own key fails the "
    "certificate check. Each STRIDE threat maps to a check. The residual risk is a compromised build "
    "server that signs malware legitimately; the log makes that visible, and SLSA provenance is the "
    "future fix. (~45 s)")
stride = [["STRIDE", "Threat", "SignGate control", "Check"],
          ["Spoofing", "Fake 'Codecov' key", "Cert must chain to pinned root", "1"],
          ["Tampering", "Edit uploader (2021 attack)", "SHA-256 bound by Ed25519", "3, 4"],
          ["Tampering", "Replay old signed version", "Signed version ≥ last seen", "5"],
          ["Repudiation", "Hidden / denied release", "Append-only signed log", "6"],
          ["Info. disclosure", "Read all CI secrets", "Env allowlist", "exec"],
          ["DoS", "Corrupt files to break builds", "Fail-closed + mirrors", "-"],
          ["Elevation", "Stolen release key", "Log + monitor + CRL", "2, 6"]]
table(s, 0.5, 1.55, 9.0, [1.75, 2.75, 3.5, 1.0], stride, size=13, row_h=0.47)
card(s, 0.5, 5.5, 9.0, 1.2, "green", border=GREEN)
text(s, 0.7, 5.55, 8.6, 1.1, [[("Original attack vector: ", True),
     ("attacker has bucket write access, no signing key → hash mismatch (check 4). Also rewrites the hash → "
      "signature invalid (check 3). Blocked before a single line runs. Residual risk: compromised build server (future: SLSA).", False)]],
     size=13.5, anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------------------ 14 results
s = new_slide("Results: System Testing",
    "We attacked our own system. The legitimate release is allowed. The exact 2021 attack is blocked "
    "at the hash check, and each stronger variant is blocked at a different check, so every layer "
    "does something. In the assume-breach test, where we run the malicious script anyway, the "
    "allowlist cut exposed secrets from 8 to 1. Verification costs about half a millisecond. (~50 s)")
rows = [["ID", "Scenario", "Expected", "Result", "Stopped at"]]
for r in RESULTS["scenarios"]:
    rows.append([r["id"], r["scenario"].replace("Codecov 2021 attack: ", "2021 attack: "),
                 "ALLOW" if r["id"] == "S1" else "BLOCK", r["result"], r["stopped_by"]])
table(s, 0.5, 1.5, 9.0, [0.5, 4.1, 1.05, 1.15, 2.2], rows, size=11.5, row_h=0.4,
      highlight=lambda ri, ci, v: (GREEN if v == "ALLOWED" else RED if v == "BLOCKED" else None) if ci == 3 else None)
picture(s, "exposure.png", 0.55, 4.95, w=5.6)
for i, (big, small) in enumerate([(f"{RESULTS['verify_ms']} ms", "avg verification"), ("13 / 13", "unit tests pass")]):
    x = 6.55 + i * 1.5
    card(s, x, 4.95, 1.4, 1.65, "navy")
    text(s, x, 5.1, 1.4, 0.7, [[(big, True)]], size=19, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, x, 5.8, 1.4, 0.7, [small], size=12, color=MUTED, align=PP_ALIGN.CENTER)

# ------------------------------------------------------------------ 15 demo output
s = new_slide("Results: Demo Output",
    "This is the actual console output of our attack simulation. Green PASS and red FAIL show "
    "exactly which check stopped each attack. The monitor also flags a tampered file immediately; in "
    "2021 that took two months. If time permits, we run this live: python demo/attack_demo.py. (~20 s)")
picture(s, "term1.png", 0.5, 1.6, w=4.4, h=5.1)
picture(s, "term2.png", 5.1, 1.6, w=4.4, h=5.1)

# ------------------------------------------------------------------ 16 innovation
s = new_slide("Innovation and Application Relevance",
    "After the incident Codecov published a signed uploader, but verification stayed optional and "
    "manual. Our contribution is not a new algorithm. It combines standard primitives into an "
    "enforced gate that also handles key theft, rollback and blast radius. It matches SLSA, NIST SSDF "
    "and Executive Order 14028, and costs about a dollar a month for a KMS key. The same pattern "
    "applies to GitHub Actions, npm and container images, as the 2025 tj-actions incident "
    "showed. (~45 s)")
cmp_rows = [["Aspect", "Standard fix (signed + SHA)", "SignGate"],
            ["Verification", "Optional, manual", "Enforced, fail-closed"],
            ["Trust anchor", "Fetched from same site", "Pinned in customer repo"],
            ["Leaked signing key", "Every user must update", "Root-signed CRL, automatic"],
            ["Silent key misuse", "Undetectable", "Transparency log + monitor"],
            ["Rollback", "Not addressed", "Signed version + state"],
            ["If malicious code runs", "All secrets exposed", "1 of 8 secrets (allowlist)"],
            ["Detection time", "~2 months", "Minutes"]]
table(s, 0.5, 1.55, 9.0, [2.4, 3.1, 3.5], cmp_rows, size=13, row_h=0.43,
      highlight=lambda ri, ci, v: GREEN if ci == 2 else None)
rel = [("Industry", "SLSA · NIST SSDF · EO 14028"), ("Cost", "Open source; KMS ≈ US$1/month"),
       ("Feasibility", "1 CI line; ~0.6 ms per build")]
for i, (h, d) in enumerate(rel):
    x = 0.5 + i * 3.05
    card(s, x, 5.2, 2.9, 1.5, "navy")
    text(s, x + 0.12, 5.3, 2.66, 0.45, [[(h, True)]], size=16, color=NAVY)
    text(s, x + 0.12, 5.78, 2.66, 0.85, [d], size=13.5)

# ------------------------------------------------------------------ 17 conclusion
s = new_slide("Conclusion",
    "To conclude: Codecov was an integrity and authenticity failure that became a confidentiality "
    "breach. SignGate fixes that with standard cryptography and good key management, and our tests "
    "show it blocks the real attack and stronger variants cheaply. Next steps are Sigstore keyless "
    "signing, SLSA provenance and packaging it as a GitHub Action. (~30 s)")
concl = [[("Root cause: ", True), ("integrity + authenticity failure (unsigned uploader, leaked static key) → confidentiality breach.", False)],
         [("SignGate: ", True), ("SHA-256 + Ed25519 signatures, root-certified rotatable keys with CRL, transparency log, "
                                 "anti-rollback and least-privilege execution in one fail-closed gate.", False)],
         [("Evidence: ", True), ("7/7 scenarios correct, 2021 attack blocked before execution, secrets exposed 8 → 1, ~0.6 ms.", False)],
         [("Lesson: ", True), ("trust keys, not locations. Verify before you execute.", False)]]
text(s, 0.5, 1.6, 9.0, 3.7, concl, size=17, bullets=True, spacing=12)
card(s, 0.5, 5.4, 9.0, 1.3, "navy")
text(s, 0.7, 5.45, 8.6, 1.2, [[("Future work: ", True),
     ("Sigstore keyless (OIDC) signing · SLSA build provenance for compromised build servers · "
      "witness co-signed log · package as a reusable GitHub Action", False)]], size=14.5, anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------------------ 18 references
s = new_slide("References (IEEE)",
    "Our main sources are Codecov's own security update and post-mortem, plus Reuters and HashiCorp "
    "for impact. The standards cited are for the algorithms we used. (~5 s)")
refs = ['[1] Codecov, "Bash Uploader Security Update," Apr. 15, 2021. about.codecov.io/security-update',
        '[2] Codecov, "Post-Mortem / Root Cause Analysis (April 2021)." about.codecov.io/apr-2021-post-mortem',
        '[3] J. Menn and R. Satter, "Codecov hackers breached hundreds of restricted customer sites," Reuters, Apr. 2021.',
        '[4] HashiCorp, "HCSEC-2021-12: Codecov Security Event and HashiCorp GPG Key Exposure," Apr. 2021.',
        '[5] S. Josefsson and I. Liusvaara, "Edwards-Curve Digital Signature Algorithm (EdDSA)," RFC 8032, 2017.',
        '[6] NIST, "Secure Hash Standard (SHS)," FIPS 180-4, 2015; "Digital Signature Standard," FIPS 186-5, 2023.',
        '[7] E. Barker, "Recommendation for Key Management," NIST SP 800-57 Pt. 1 Rev. 5, 2020.',
        '[8] B. Laurie, A. Langley and E. Kasper, "Certificate Transparency," RFC 6962, 2013.',
        '[9] Z. Newman et al., "Sigstore: Software Signing for Everybody," ACM CCS, 2022.',
        '[10] OpenSSF, "SLSA v1.0," 2023; NIST SP 800-218 (SSDF), 2022; Executive Order 14028, 2021.']
text(s, 0.5, 1.55, 9.0, 5.3, refs, size=13, spacing=7, font=SERIF)

# ------------------------------------------------------------------ 19 thank you
s = new_slide("", "Thank you. We are happy to take questions or run the demo live. ")
text(s, 0.5, 2.4, 9.0, 1.2, [[("Thank You", True)]], size=48, color=NAVY, align=PP_ALIGN.CENTER, font=SERIF)
text(s, 0.5, 3.7, 9.0, 0.6, ["Questions?"], size=24, color=MUTED, align=PP_ALIGN.CENTER, font=SERIF)
card(s, 2.0, 4.8, 6.0, 0.8, "navy")
text(s, 2.0, 4.8, 6.0, 0.8, [[("Trust keys, not locations. Verify before you execute.", True)]], size=16,
     color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(sys.argv[2])
print("saved", sys.argv[2], "slides:", len(prs.slides))
